#!/usr/bin/env python3
"""Genera presentaciones adaptativas y verificables para las sesiones de MiniJarvis.

La guía docente es la fuente principal. La ficha del alumnado complementa seguridad,
consignas y cierre cuando la guía no contiene esos elementos. La generación admite
validación, simulación, selección de sesiones y publicación con copia de seguridad.
"""

from __future__ import annotations

import argparse
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
STUDENT_ROOT = ROOT / "01-ALUMNADO" / "03-SESIONES"
TEACHER_ROOT = ROOT / "02-PROFESORADO" / "02-SESIONES"
OUTPUT_ROOT = ROOT / "02-PROFESORADO" / "03-PRESENTACIONES" / "POR-SESION"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = RGBColor(28, 35, 48)
MUTED = RGBColor(88, 98, 112)
PAPER = RGBColor(249, 247, 242)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(229, 233, 239)
DARK = RGBColor(18, 24, 36)
GREEN = RGBColor(36, 126, 92)
RED = RGBColor(177, 58, 58)
AMBER = RGBColor(192, 112, 32)

COLORS = {
    "h0": RGBColor(204, 79, 62),
    "h1": RGBColor(47, 91, 168),
    "h2": RGBColor(20, 124, 136),
    "h3": RGBColor(53, 129, 84),
    "c1": RGBColor(111, 79, 154),
    "h4": RGBColor(105, 75, 157),
    "h5": RGBColor(200, 112, 34),
    "c2": RGBColor(111, 79, 154),
    "h6": RGBColor(33, 89, 122),
    "h7": RGBColor(156, 57, 117),
    "hf": RGBColor(109, 87, 37),
}

@dataclass
class TimelineBlock:
    time: str
    action: str

@dataclass
class Session:
    number: str
    folder: str
    topic: str
    hito: str
    duration: str
    moment: str
    grouping: str
    objective: str
    evidence: str
    materials: list[str]
    timeline: list[TimelineBlock]
    key_concepts: list[str]
    example: list[str]
    activity: list[str]
    checklist: list[str]
    safety: list[str]
    closure: list[str]
    close_question: str
    teacher_source: Path
    student_source: Path

@dataclass
class SlideSpec:
    kind: str
    title: str
    items: list[str] = field(default_factory=list)
    subtitle: str = ""

@dataclass
class ValidationReport:
    session: str
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

GENERIC_MARKERS = {
    "Comprender y aplicar el objetivo de la sesión.",
    "Una evidencia comprobable y defendible.",
    "Sigue la consigna indicada en clase.",
    "¿Qué has hecho y cómo sabes que funciona?",
}

def clean(text: str) -> str:
    text = re.sub(r"\[([^]]+)]\([^)]*\)", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"<[^>]+>", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text.replace(".,", ",").replace(";,", ";")

def normalize_projectable_text(text: str) -> str:
    text = re.sub(r"docs/depuracion-h2(?!\.md)", "docs/depuracion-h2.md", text)
    text = text.replace(
        "docs/incidencia-h6.md/docs/seguridad-h6.md",
        "docs/incidencia-h6.md o docs/seguridad-h6.md",
    )
    return text.replace(
        "docs/incidencia-h6/docs/seguridad-h6",
        "docs/incidencia-h6.md o docs/seguridad-h6.md",
    )

def heading_pattern(title_pattern: str) -> str:
    return rf"(?:\d+(?:\.\d+)*[.)]?\s+)?(?:{title_pattern})"

def section(text: str, title_pattern: str, levels: tuple[int, ...] = (2, 3)) -> str:
    """Devuelve una sección Markdown aunque el título tenga prefijo numérico."""
    lines = text.splitlines()
    title_re = re.compile(rf"^{heading_pattern(title_pattern)}", flags=re.IGNORECASE)
    for index, line in enumerate(lines):
        match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if not match:
            continue
        level = len(match.group(1))
        if level not in levels or not title_re.search(match.group(2)):
            continue
        body: list[str] = []
        for candidate in lines[index + 1 :]:
            next_heading = re.match(r"^(#{1,6})\s+", candidate)
            if next_heading and len(next_heading.group(1)) <= level:
                break
            body.append(candidate)
        return "\n".join(body).strip()
    return ""

def list_items(text: str, limit: int | None = None) -> list[str]:
    items: list[str] = []
    for line in text.splitlines():
        match = re.match(r"\s*(?:\d+\.|[-*]\s*(?:\[[ xX]\])?)\s+(.+)", line)
        if not match:
            continue
        value = clean(match.group(1))
        if value and value not in items:
            items.append(value)
        if limit is not None and len(items) >= limit:
            break
    return items

def is_noise(value: str) -> bool:
    return bool(
        re.search(
            r"^(?:Guion breve sugerido|Hoy necesitamos comprender y practicar lo justo|Producto o evidencia que debe quedar)",
            value,
            flags=re.IGNORECASE,
        )
    )

def meaningful_lines(text: str, limit: int | None = None) -> list[str]:
    result: list[str] = []
    in_code = False
    for raw in text.splitlines():
        line = raw.strip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if not line or line.startswith("|") or re.match(r"^[-:|. ]+$", line):
            continue
        if line.startswith("#"):
            value = clean(line.lstrip("# "))
            value = re.sub(r"^\d+(?:\.\d+)*[.)]?\s*", "", value)
        elif re.match(r"^(?:\d+\.|[-*])\s+", line):
            value = clean(re.sub(r"^(?:\d+\.|[-*])\s+", "", line))
        elif in_code or not line.startswith("#"):
            value = clean(line.lstrip("> "))
        else:
            continue
        if value and not is_noise(value) and not set(value) <= {".", "_"} and value not in result:
            result.append(value)
        if limit is not None and len(result) >= limit:
            break
    return result

def semantic_items(text: str, limit: int | None = None) -> list[str]:
    """Agrupa introducciones, listas y filas de tabla en ideas proyectables."""
    items: list[str] = []
    pending_label = ""
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("```") or re.fullmatch(r"[-:|. ]+", line):
            continue
        if line.startswith("#"):
            pending_label = clean(re.sub(r"^#+\s+(?:\d+(?:\.\d+)*[.)]?\s*)?", "", line))
            continue
        if line.startswith("|"):
            cells = [clean(cell) for cell in line.strip("|").split("|")]
            if len(cells) < 2 or any(re.fullmatch(r"[-: ]+", cell) for cell in cells):
                continue
            if cells[0].lower() in {"función inicial", "tendencia hada", "campo", "dato"}:
                continue
            value = " — ".join(cell for cell in cells if cell)
            if value and value not in items:
                items.append(value)
            pending_label = ""
        else:
            bullet = re.match(r"^(?:\d+\.|[-*])\s+(.+)", line)
            value = clean(bullet.group(1) if bullet else line.lstrip("> "))
            if not value or is_noise(value):
                continue
            if bullet and items and items[-1].endswith(":"):
                items[-1] = f"{items[-1]} {value}"
            elif bullet and items and ": " in items[-1] and len(items[-1]) < 250:
                items[-1] = f"{items[-1].rstrip(' ;,.')}, {value}"
            else:
                if pending_label and not value.lower().startswith(pending_label.lower()):
                    value = f"{pending_label}: {value}"
                if value not in items:
                    items.append(value)
                pending_label = ""
        if limit is not None and len(items) >= limit:
            break
    return items

def operational_details(text: str, limit: int = 3) -> list[str]:
    """Conserva el contenido anunciado por etiquetas como ``Di:`` o ``Pregunta:``."""
    items = semantic_items(text)
    details: list[str] = []
    index = 0
    while index < len(items) and len(details) < limit:
        item = items[index]
        if item.endswith(":") and index + 1 < len(items):
            announced = items[index + 1].replace("…", "____")
            details.append(f"{item} {announced}")
            index += 2
            continue
        details.append(item.rstrip(":").replace("…", "____"))
        index += 1
    return details

def metadata_table(text: str) -> dict[str, str]:
    values: dict[str, str] = {}
    for line in text.splitlines():
        if not line.strip().startswith("|"):
            continue
        cells = [clean(cell) for cell in line.strip().strip("|").split("|")]
        if len(cells) < 2 or not cells[0] or re.fullmatch(r"[-: ]+", cells[0]):
            continue
        if cells[0].lower() in {"dato", "campo", "hoy vas a…", "hoy vais a…"}:
            continue
        values.setdefault(cells[0].lower(), cells[1])
    return values

def first_student_outcome(text: str) -> tuple[str, str]:
    lines = text.splitlines()
    for index, line in enumerate(lines):
        normalized = clean(line).lower()
        if ("hoy vas" in normalized or "hoy vais" in normalized) and "debe quedar" in normalized:
            for candidate in lines[index + 2 : index + 6]:
                if candidate.strip().startswith("|"):
                    cells = [clean(cell) for cell in candidate.strip().strip("|").split("|")]
                    if len(cells) >= 2:
                        return cells[0], cells[1]
    return "", ""

def value_for(meta: dict[str, str], *keys: str) -> str:
    for key in keys:
        value = meta.get(key.lower(), "")
        if value:
            return value
    return ""

def parse_markdown_timeline(text: str) -> list[TimelineBlock]:
    body = section(text, r"Secuencia de aula")
    timeline: list[TimelineBlock] = []
    if body:
        for line in body.splitlines():
            if not line.strip().startswith("|"):
                continue
            cells = [clean(cell) for cell in line.strip().strip("|").split("|")]
            if len(cells) >= 2 and re.search(r"\d", cells[0]) and not re.fullmatch(r"[-: ]+", cells[0]):
                timeline.append(TimelineBlock(cells[0], cells[1]))
    if timeline:
        return timeline

    headings = list(
        re.finditer(
            r"^###\s+Tramo\s+\d+\s+[—-]\s*(.+?)(?:,|\s)[ ]*(\d+\s*[–-]\s*\d+)\s+minutos.*?$",
            text,
            flags=re.MULTILINE | re.IGNORECASE,
        )
    )
    for index, match in enumerate(headings):
        start = match.end()
        end = headings[index + 1].start() if index + 1 < len(headings) else len(text)
        block = text[start:end]
        details = operational_details(block)
        action = clean(match.group(1))
        if details:
            action = f"{action}: {' · '.join(details)}"
        timeline.append(TimelineBlock(f"{clean(match.group(2))} min", action))
    if timeline:
        return timeline

    minute_headings = list(
        re.finditer(
            r"^###\s+Minutos\s+(\d+\s*[–-]\s*\d+)\s+[—-]\s*(.+?)\s*$",
            text,
            flags=re.MULTILINE | re.IGNORECASE,
        )
    )
    for index, match in enumerate(minute_headings):
        start = match.end()
        end = minute_headings[index + 1].start() if index + 1 < len(minute_headings) else len(text)
        details = operational_details(text[start:end])
        action = clean(match.group(2))
        if details:
            action = f"{action}: {' · '.join(details)}"
        timeline.append(TimelineBlock(f"{clean(match.group(1))} min", action))
    return timeline

def extract_question(teacher_text: str, student_text: str) -> str:
    for source, heading in (
        (student_text, r"Cierre(?: individual)?"),
        (teacher_text, r"Comprobación final"),
    ):
        body = section(source, heading)
        bold = re.search(r"\*\*(.+?)\*\*", body, flags=re.DOTALL)
        if bold:
            question = clean(bold.group(1))
            question = re.sub(r"^(?:Pregunta(?: de control| final)?):\s*", "", question, flags=re.I)
            question = re.sub(r"\s*#{2,}.*$", "", question)
            return question
        lines = meaningful_lines(body, 2)
        if lines:
            return lines[0]
    return ""

def teacher_for_student(student_path: Path) -> Path:
    relative = student_path.relative_to(STUDENT_ROOT)
    return TEACHER_ROOT / relative.with_name(relative.name.replace("-alumnado.md", "-docente.md"))

def student_for_teacher(teacher_path: Path) -> Path:
    relative = teacher_path.relative_to(TEACHER_ROOT)
    return STUDENT_ROOT / relative.with_name(relative.name.replace("-docente.md", "-alumnado.md"))

def parse_session(teacher_path: Path, student_path: Path | None = None) -> Session:
    """Analiza una sesión usando la guía docente como fuente de autoridad."""
    if student_path is None:
        student_path = student_for_teacher(teacher_path)
    teacher = teacher_path.read_text(encoding="utf-8")
    student = student_path.read_text(encoding="utf-8") if student_path.exists() else ""
    meta = metadata_table(teacher)
    student_objective, student_evidence = first_student_outcome(student)

    number_match = re.search(r"S(\d{3})", teacher_path.name)
    number = number_match.group(1) if number_match else "000"
    headings = re.findall(r"^##\s+(.+)$", teacher, flags=re.MULTILINE)
    topic = clean(headings[0]) if headings else teacher_path.stem

    objective = value_for(meta, "Producto principal", "Resultado observable", "Resultado de hoy")
    if not objective:
        objective = student_objective
    evidence = value_for(meta, "Evidencia individual", "Evidencia mínima")
    if not evidence:
        evidence = student_evidence

    duration = value_for(meta, "Duración", "Duración prevista")
    hito = value_for(meta, "Hito") or teacher_path.parent.name.upper()
    moment = value_for(
        meta,
        "Fase HEXA",
        "Fase HEXA del hito",
        "Momento HEXA",
        "Momento HEXA del hito",
    )
    grouping = value_for(meta, "Agrupamiento")

    materials_body = section(teacher, r"Material imprescindible") or section(teacher, r"Material", levels=(3,))
    materials = list_items(materials_body)
    if not materials:
        materials = list_items(section(student, r"Material(?: que necesitas| disponible)"))

    key_bodies: list[str] = []
    for key_heading in (
        r"Qué debes explicar",
        r"Distinción imprescindible",
        r"Qué explicas tú y qué descubre el alumnado",
        r"Qué debes dominar antes de explicarlo",
        r"Reglas que deben explicarse sin ambigüedad",
        r"Marco de decisión sobre los equipos",
    ):
        body = section(teacher, key_heading)
        if body and body not in key_bodies:
            key_bodies.append(body)
    key_concepts = semantic_items("\n\n".join(key_bodies), 24)

    example_body = section(teacher, r"Ejemplo o demostración preparada")
    if not example_body:
        example_body = section(teacher, r"Preparación de la pizarra", levels=(3,))
    example = semantic_items(example_body, 8)

    timeline = parse_markdown_timeline(teacher)

    activity_bodies = [section(teacher, r"Consigna que se entrega al alumnado")]
    for activity_heading in (
        r"Trabajo de hoy",
        r"Acuerdo de funciones",
        r"Regla de participación",
        r"Reto que prepararemos",
        r"Backlog inicial de la torre",
        r"Definición de terminado y bloqueo",
        r"Defensa breve del diseño del equipo",
        r"Cierre individual",
    ):
        body = section(student, activity_heading)
        if body:
            activity_bodies.append(body)
    activity = semantic_items("\n\n".join(body for body in activity_bodies if body), 36)
    if not activity:
        activity = [block.action for block in timeline if re.search(r"equipo|completar|crear|decidir|trabajo|backlog|defensa", block.action, re.I)]

    observe = list_items(section(teacher, r"Qué observar mientras trabajan"))
    checklist = list_items(section(teacher, r"Criterios? (?:de cierre|para considerar cerrada la sesión)"))
    if not checklist:
        checklist = list_items(section(student, r"Evidencia mínima antes de salir"))
    if not checklist:
        checklist = observe

    safety = list_items(section(student, r"Seguridad y uso de IA"))
    if not safety:
        safety = [item for item in materials if re.search(r"datos|contraseñas|tokens|claves|personales", item, re.I)]
    if not safety:
        candidate_rules = list_items(section(student, r"Reglas de esta sesión"))
        safety = [item for item in candidate_rules if re.search(r"publicar|puntuaciones|datos|consentimiento|particip", item, re.I)]

    closure_body = section(teacher, r"Comprobación final") or section(teacher, r"Criterios? de cierre")
    closure = meaningful_lines(closure_body, 10)
    close_question = extract_question(teacher, student)

    return Session(
        number=number,
        folder=teacher_path.parent.name,
        topic=topic,
        hito=hito,
        duration=duration,
        moment=moment,
        grouping=grouping,
        objective=objective,
        evidence=evidence,
        materials=materials,
        timeline=timeline,
        key_concepts=key_concepts,
        example=example,
        activity=activity,
        checklist=checklist,
        safety=safety,
        closure=closure,
        close_question=close_question,
        teacher_source=teacher_path,
        student_source=student_path,
    )

def split_long_text(text: str, max_chars: int = 170) -> list[str]:
    """Divide texto largo sin recortarlo ni introducir puntos suspensivos."""
    text = clean(text)
    if len(text) <= max_chars:
        return [text] if text else []
    parts = [clean(part) for part in re.split(r"(?<=[.;:])\s+|\s+[—-]\s+", text) if clean(part)]
    if len(parts) == 1:
        parts = [clean(part) for part in re.split(r"(?<=,)\s+", text) if clean(part)]
    result: list[str] = []
    current = ""
    for part in parts:
        candidate = f"{current} {part}".strip()
        if current and len(candidate) > max_chars:
            result.append(current)
            current = part
        else:
            current = candidate
    if current:
        result.append(current)
    return result

def expanded_items(items: list[str], max_chars: int = 170) -> list[str]:
    expanded: list[str] = []
    for item in items:
        expanded.extend(split_long_text(item, max_chars))
    return expanded

def chunk_items(items: list[str], max_items: int, max_chars: int) -> list[list[str]]:
    chunks: list[list[str]] = []
    current: list[str] = []
    size = 0
    for item in items:
        if current and (len(current) >= max_items or size + len(item) > max_chars):
            chunks.append(current)
            current = []
            size = 0
        current.append(item)
        size += len(item)
    if current:
        chunks.append(current)
    if len(chunks) > 1 and len(chunks[-1]) == 1 and len(chunks[-2]) > 2:
        chunks[-1].insert(0, chunks[-2].pop())
    return chunks

def timeline_display_item(block: TimelineBlock) -> str:
    """Resume la agenda sin repetir el desarrollo detallado ni dejar etiquetas abiertas."""
    action = clean(block.action)
    title, separator, detail = action.partition(": ")
    if not separator:
        return f"{block.time} — {action.rstrip(':')}"
    detail = re.sub(r"^(?:Di|Pregunta|Consigna literal|Anuncia|Presenta):\s*", "", detail, flags=re.I)
    first_sentence = re.split(r"(?<=[.!?])\s+|\s+·\s+", detail, maxsplit=1)[0]
    summary = f"{title}: {first_sentence}".rstrip(":")
    return f"{block.time} — {summary}"

def projectable_activity_items(session: Session) -> list[str]:
    """Condensa actividades densas en fases completas y proyectables."""
    corpus = " ".join([session.topic, *session.activity]).lower()
    if "hada" in corpus and "torre" in corpus:
        return [
            "Acuerdo de funciones — responsable, sustituto, razón basada en aporte o aprendizaje y conducta observable.",
            "Regla de participación — impedir que una sola persona diseñe o construya toda la torre.",
            "Ciclo 1 — 6 hojas A4 reutilizadas, tijeras y regla, sin conectores.",
            "Ciclo 2 — solicitar material adicional solo tras justificar problema, cambio previsto y prueba esperada.",
            "Reto — torre autoportante dentro de 25 × 25 cm; carga aproximada de 50 g durante 60 segundos.",
            "Backlog — al menos seis tareas comprobables, con prioridad y criterio de Hecho.",
            "Definición de terminado y bloqueo probable — indicar cómo se harán visibles.",
            "Defensa — explicar capacidad a compensar, función, conducta observable y criterio de terminado.",
            "Cierre individual — Mi responsabilidad inicial; conducta observable; tarea compartida; condición para revisar la función.",
        ]
    if "breakpoint" in corpus and "username" in corpus:
        return [
            "Coloca el breakpoint después de leer command.",
            "Inicia el programa en modo Debug e introduce un dato ficticio.",
            "Avanza paso a paso hasta antes del condicional.",
            "Observa command, running y userName y registra sus valores.",
            "Guarda una captura o descripción comprobable en docs/depuracion-h2.md.",
        ]
    if "command" in corpus and "sobreingeniería" in corpus:
        return [
            "Nombra el problema de diseño observado en el proyecto.",
            "Describe una alternativa simple antes de introducir un patrón.",
            "Explica la solución utilizada y localiza el código real que la sostiene.",
            "Formula una semejanza concreta con Command: cada Tool encapsula una acción ejecutable.",
            "Formula una diferencia: no hacen falta invocadores, receptores o fábricas si no aportan valor.",
            "Identifica un riesgo comprobable de sobreingeniería.",
            "Decide mantener, simplificar o descartar y registra una justificación defendible.",
        ]
    if "memory" in corpus and ("excepción" in corpus or "error" in corpus):
        return [
            "Provoca de forma controlada un fallo de carga o guardado.",
            "Crea una excepción propia o justifica por qué empleas una estándar; revisa throws.",
            "Controla carga y guardado sin ocultar el error al usuario.",
            "Impide que Memory guarde recuerdos nulos o vacíos y que exponga su lista interna modificable.",
            "Prueba el error controlado y documenta la decisión en docs/incidencia-h6.md o docs/seguridad-h6.md.",
        ]
    if "defensa final" in corpus and "retrospectiva" in corpus:
        return [
            "Prepara una afirmación y localiza el artefacto que la respalda.",
            "Defiende individualmente el producto con una prueba observable y una decisión explicada.",
            "Responde una pregunta vinculando producto, proceso y aprendizaje.",
            "Completa la autoevaluación y la retrospectiva con una evidencia concreta.",
            "Formula y entrega un plan personal de mejora con un siguiente paso verificable.",
        ]
    boilerplate = (
        "Lee el objetivo",
        "Atiende el ejemplo",
        "Atiende al ejemplo",
        "Realiza esta tarea",
        "no basta con decir",
        "Guarda o entrega la evidencia",
    )
    specific = [item for item in session.activity if not item.startswith(boilerplate)]
    if len(specific) == 1:
        specific.append(f"Comprueba y registra la evidencia: {session.evidence}")
    return specific

def projectable_checklist_items(session: Session) -> list[str]:
    if session.folder == "h1":
        student = session.student_source.read_text(encoding="utf-8")
        evidence = list_items(section(student, r"Evidencia única antes de salir"))
        if evidence:
            return evidence

    corpus = " ".join([session.topic, session.evidence, *session.activity]).lower()
    if "breakpoint" in corpus and "username" in corpus:
        return [
            "El breakpoint está situado después de leer command.",
            "La ejecución se ha iniciado en modo Debug con datos ficticios.",
            "Se han observado y registrado command, running y userName.",
            "docs/depuracion-h2.md contiene una captura o descripción comprobable.",
            "Puedo explicar qué reveló la depuración antes del condicional.",
        ]
    if "memory" in corpus and ("excepción" in corpus or "error" in corpus):
        return [
            "El error de carga o guardado se reproduce de forma controlada.",
            "La excepción elegida y el uso de throws están justificados.",
            "Memory rechaza recuerdos nulos o vacíos y protege su lista interna.",
            "Existe una prueba observable del estado seguro.",
            "La decisión está documentada en docs/incidencia-h6.md o docs/seguridad-h6.md.",
        ]
    if "defensa final" in corpus and "retrospectiva" in corpus:
        return [
            "Defensa: producto funcional, prueba observable y decisión explicada.",
            "Trazabilidad: puedo localizar los cambios y su justificación.",
            "Aportación individual: identifico qué hice y qué aprendí.",
            "Autoevaluación y retrospectiva: aportan una evidencia concreta.",
            "Plan personal de mejora: incluye un siguiente paso verificable.",
        ]
    return session.checklist

def projectable_close_items(session: Session) -> list[str]:
    corpus = " ".join([session.topic, *session.activity]).lower()
    if "hada" in corpus and "torre" in corpus:
        return [
            "Mi responsabilidad inicial",
            "La conducta observable que demostrará que la cumplo",
            "Una tarea compartida en la que también participaré",
            "La condición para revisar la función después de la torre",
        ]
    if session.close_question:
        return [session.close_question]
    return expanded_items(session.closure, 175)

def disciplinary_visuals(session: Session) -> list[SlideSpec]:
    """Selecciona apoyos visuales específicos a partir del contenido didáctico."""
    corpus = " ".join(
        [
            session.topic,
            session.objective,
            session.evidence,
            *session.key_concepts,
            *session.example,
            *session.activity,
        ]
    )
    lowered = corpus.lower()
    visuals: list[SlideSpec] = []

    if "hada" in lowered and "funcion" in lowered:
        visuals.append(
            SlideSpec(
                "relationship",
                "HADA orienta; las funciones se acuerdan",
                [
                    "Gestor — natural: valor, backlog y tiempo · practicar: escucha y revisión",
                    "Colaborador — natural: facilitación y mediación · practicar: priorización y crítica",
                    "Desarrollador — natural: ideas y prototipado · practicar: cierre, prueba y documentación",
                    "Analista — natural: calidad y prueba · practicar: experimentación y comunicación",
                    "La tabla orienta una conversación; no asigna funciones automáticamente.",
                ],
                "Hipótesis para conversar, no etiquetas para encasillar",
            )
        )

    if "breakpoint" in lowered and all(name.lower() in lowered for name in ("command", "running", "username")):
        visuals.append(
            SlideSpec(
                "debugger",
                "Depurar en IntelliJ: observar antes de decidir",
                [
                    "1. Coloca el breakpoint después de leer command",
                    "2. Inicia el programa en modo Debug",
                    "3. Avanza paso a paso hasta antes del condicional",
                    "4. Observa command, running y userName en Variables",
                    "5. Guarda captura o descripción en docs/depuracion-h2.md",
                ],
                "Breakpoint → Debug → paso a paso → Variables → evidencia",
            )
        )

    if "command" in lowered and "sobreingeniería" in lowered:
        visuals.append(
            SlideSpec(
                "pattern",
                "Command simplificado: problema, semejanza y límite",
                [
                    "Problema real — evitar que cada acción engorde un condicional central",
                    "Una semejanza con Command — cada Tool encapsula una acción ejecutable",
                    "Una diferencia — no hacen falta invocadores, receptores o fábricas si no aportan valor",
                    "Decisión — mantener, simplificar o descartar según el coste de sobreingeniería",
                ],
            )
        )

    if "throws" in lowered and "invariante" in lowered:
        visuals.append(
            SlideSpec(
                "exception_flow",
                "Del fallo de fichero a un estado seguro",
                [
                    "Fallo de carga o guardado",
                    "Lanzar o propagar con throws",
                    "Capturar o traducir a MemoryStorageException / PersistenceException",
                    "Proteger la invariante: Memory no guarda recuerdos nulos o vacíos",
                    "Comprobar el error controlado y documentar en docs/incidencia-h6.md o docs/seguridad-h6.md",
                ],
            )
        )

    if "defensa final" in lowered and "retrospectiva" in lowered:
        visuals.extend(
            [
                SlideSpec(
                    "defense",
                    "Qué hace defendible la evidencia final",
                    [
                        "Producto funcional",
                        "Prueba observable",
                        "Decisión explicada",
                        "Trazabilidad de cambios",
                        "Aportación individual",
                        "Plan personal de mejora",
                    ],
                    "Construir · probar · documentar · defender",
                ),
                SlideSpec(
                    "microdefense",
                    "Microdefensa ejemplificada",
                    [
                        "Afirmación — qué funciona y para quién",
                        "Artefacto — dónde puede verse en el proyecto",
                        "Prueba — qué ejecución o resultado lo demuestra",
                        "Decisión — qué alternativa se eligió y por qué",
                        "Dificultad y aprendizaje — qué cambiarías ahora",
                    ],
                ),
            ]
        )
    return visuals

def plan_slides(session: Session) -> list[SlideSpec]:
    """Construye un plan variable según el contenido real de la sesión."""
    is_h1 = session.folder == "h1"
    title_subtitle = session.hito
    if is_h1 and session.moment:
        title_subtitle = f"{session.hito} · Fase HEXA: {session.moment}"
    elif session.duration:
        title_subtitle = f"{session.hito} · {session.duration}"
    slides = [SlideSpec("title", session.topic, [], title_subtitle)]

    outcome_source = [f"Objetivo: {session.objective}"]
    if is_h1:
        outcome_source.append(
            "Al terminar: señala la evidencia, explica qué demuestra y reproduce la prueba."
        )
    else:
        outcome_source.append(f"Evidencia: {session.evidence}")
    outcome_items = expanded_items(outcome_source, 190)
    for index, chunk in enumerate(chunk_items(outcome_items, 4, 520), 1):
        suffix = "" if index == 1 else f" · {index}"
        slides.append(SlideSpec("outcome", f"Qué debe conseguirse{suffix}", chunk, session.moment))

    if not is_h1:
        timeline_items = [timeline_display_item(block) for block in session.timeline]
        timeline_chunk_size = 6 if len(timeline_items) > 7 else 7
        for index, chunk in enumerate(chunk_items(timeline_items, timeline_chunk_size, 1200), 1):
            suffix = "" if len(timeline_items) <= timeline_chunk_size else f" · tramo {index}"
            slides.append(SlideSpec("timeline", f"Secuencia de aula{suffix}", chunk))

    visuals = disciplinary_visuals(session)
    visual_kinds = {spec.kind for spec in visuals}
    replaces_sparse_explanation = bool(visual_kinds & {"debugger", "exception_flow", "defense"})

    concepts = expanded_items(session.key_concepts, 210)
    examples = expanded_items(session.example, 170)
    sparse_focus = not visuals and 1 <= len(concepts) <= 2 and 1 <= len(examples) <= 2
    if sparse_focus:
        slides.append(
            SlideSpec(
                "focus",
                "Idea clave y puesta en práctica",
                [*concepts, *examples],
                str(len(concepts)),
            )
        )
    elif not is_h1 and not (replaces_sparse_explanation and len(concepts) <= 2):
        for index, chunk in enumerate(chunk_items(concepts, 7, 900), 1):
            suffix = "" if len(concepts) <= 7 else f" · {index}"
            slides.append(SlideSpec("concepts", f"Claves para explicar{suffix}", chunk))

    slides.extend(visuals)

    if not sparse_focus and not (replaces_sparse_explanation and len(examples) <= 1):
        for index, chunk in enumerate(chunk_items(examples, 5, 650), 1):
            suffix = "" if len(examples) <= 5 else f" · {index}"
            slides.append(SlideSpec("example", f"Ejemplo o demostración{suffix}", chunk))

    activity = expanded_items(projectable_activity_items(session), 190)
    for index, chunk in enumerate(chunk_items(activity, 5, 900), 1):
        suffix = "" if len(activity) <= 5 else f" · {index}"
        start_step = str((index - 1) * 5 + 1)
        slides.append(SlideSpec("activity", f"Actividad central{suffix}", chunk, start_step))

    checklist = expanded_items(projectable_checklist_items(session), 180)
    for index, chunk in enumerate(chunk_items(checklist, 7, 950), 1):
        suffix = "" if len(checklist) <= 7 else f" · {index}"
        slides.append(SlideSpec("evidence", f"Evidencia y comprobación{suffix}", chunk))

    safety = expanded_items(session.safety, 170)
    for index, chunk in enumerate(chunk_items(safety, 6, 680), 1):
        suffix = "" if len(safety) <= 6 else f" · {index}"
        slides.append(SlideSpec("safety", f"Seguridad y uso responsable{suffix}", chunk))

    close_items = projectable_close_items(session)
    for index, chunk in enumerate(chunk_items(close_items, 4, 600), 1):
        suffix = "" if len(close_items) <= 4 else f" · {index}"
        slides.append(SlideSpec("closure", f"Cierre y defensa{suffix}", chunk))

    return [
        SlideSpec(
            slide.kind,
            normalize_projectable_text(slide.title),
            [normalize_projectable_text(item) for item in slide.items],
            normalize_projectable_text(slide.subtitle),
        )
        for slide in slides
    ]

def validate_session(session: Session) -> ValidationReport:
    report = ValidationReport(session.number)
    required = {
        "tema": session.topic,
        "hito": session.hito,
        "duración": session.duration,
        "objetivo": session.objective,
        "evidencia": session.evidence,
        "secuencia": session.timeline,
        "conceptos": session.key_concepts,
        "actividad": session.activity,
        "cierre": session.close_question or session.closure,
    }
    for name, value in required.items():
        if not value:
            report.errors.append(f"Falta {name}")
    for marker in GENERIC_MARKERS:
        if marker in {session.objective, session.evidence, session.close_question}:
            report.errors.append(f"Texto genérico no permitido: {marker}")
    plan = plan_slides(session)
    if any("…" in text for slide in plan for text in [slide.title, slide.subtitle, *slide.items]):
        report.errors.append("El plan contiene truncamientos")
    if not session.materials:
        report.warnings.append("No se extrajo material")
    if not session.safety:
        report.warnings.append("No se extrajeron reglas específicas de seguridad")
    return report

def add_full_background(slide, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

def add_text(slide, x, y, w, h, text, size=24, color=INK, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.05)
    frame.margin_top = frame.margin_bottom = Inches(0.04)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box

def body_size(items: list[str], maximum: int = 22, minimum: int = 14) -> int:
    total = sum(len(item) for item in items)
    if total > 650:
        return minimum
    if total > 480:
        return max(minimum, maximum - 5)
    if total > 320:
        return max(minimum, maximum - 3)
    return maximum

def add_bullets(slide, x, y, w, h, items: list[str], size=20, color=INK, numbered=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.04)
    for index, item in enumerate(items, start=1):
        paragraph = frame.paragraphs[0] if index == 1 else frame.add_paragraph()
        paragraph.text = f"{index}. {item}" if numbered else f"•  {item}"
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(9)
    return box

def add_footer(slide, session: Session, accent: RGBColor) -> None:
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(7.22), SLIDE_W, Inches(0.28))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, Inches(0.45), Inches(7.25), Inches(12.3), Inches(0.18), f"MiniJarvis · 1.º DAW · Sesión {session.number} · {session.hito}", 9, WHITE)

def base_slide(prs: Presentation, session: Session, title: str, kicker: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, PAPER)
    accent = COLORS.get(session.folder, RGBColor(47, 91, 168))
    side = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(0.18), SLIDE_H)
    side.fill.solid()
    side.fill.fore_color.rgb = accent
    side.line.fill.background()
    if kicker:
        add_text(slide, Inches(0.55), Inches(0.30), Inches(6.0), Inches(0.28), kicker.upper(), 10, accent, True)
    add_text(slide, Inches(0.55), Inches(0.66), Inches(12.1), Inches(0.58), title, 28, INK, True)
    add_footer(slide, session, accent)
    return slide, accent

def rounded_card(slide, x, y, w, h, fill=WHITE, line=SOFT):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = line
    return card

def render_title(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = COLORS.get(session.folder, RGBColor(47, 91, 168))
    add_full_background(slide, DARK)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(0.28), SLIDE_H)
    stripe.fill.solid(); stripe.fill.fore_color.rgb = accent; stripe.line.fill.background()
    badge = rounded_card(slide, Inches(0.72), Inches(0.62), Inches(2.05), Inches(0.62), accent, accent)
    add_text(slide, Inches(0.8), Inches(0.75), Inches(1.88), Inches(0.25), f"SESIÓN {session.number}", 14, WHITE, True, PP_ALIGN.CENTER)
    title_size = 38 if len(spec.title) < 65 else 32
    add_text(slide, Inches(0.75), Inches(1.62), Inches(11.8), Inches(1.85), spec.title, title_size, WHITE, True)
    add_text(slide, Inches(0.78), Inches(4.05), Inches(11.4), Inches(0.45), spec.subtitle, 19, RGBColor(190, 200, 214), True)
    if session.moment:
        add_text(slide, Inches(0.78), Inches(4.72), Inches(11.2), Inches(0.55), f"Fase HEXA: {session.moment}", 19, WHITE)
    if session.grouping:
        add_text(slide, Inches(0.78), Inches(5.43), Inches(11.2), Inches(0.65), session.grouping, 17, RGBColor(210, 217, 228))
    add_text(slide, Inches(0.78), Inches(6.86), Inches(11.6), Inches(0.22), "MINIJARVIS · PROGRAMACIÓN", 10, RGBColor(160, 170, 185), True)

def render_outcome(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Punto de partida")
    items = spec.items or [session.objective, session.evidence]
    columns = 2 if len(items) <= 4 else 3
    width = 11.9 / columns
    for index, item in enumerate(items):
        col = index % columns; row = index // columns
        x = 0.65 + col * width; y = 1.52 + row * 2.55
        h = 2.25 if len(items) > columns else 4.7
        rounded_card(slide, Inches(x), Inches(y), Inches(width - 0.28), Inches(h))
        add_text(slide, Inches(x + 0.25), Inches(y + 0.25), Inches(width - 0.78), Inches(h - 0.5), item, body_size([item], 21, 15), INK, index < 2)

def render_timeline(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Ritmo de la sesión")
    count = max(1, len(spec.items)); step_h = min(1.02, 5.55 / count); size = body_size(spec.items, 19, 14)
    y = 1.40
    for index, item in enumerate(spec.items, 1):
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.72), Inches(y), Inches(0.52), Inches(0.52))
        circle.fill.solid(); circle.fill.fore_color.rgb = accent; circle.line.fill.background()
        add_text(slide, Inches(0.72), Inches(y + 0.11), Inches(0.52), Inches(0.22), str(index), 13, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, Inches(1.52), Inches(y - 0.02), Inches(11.0), Inches(step_h - 0.08), item, size, INK)
        y += step_h

def render_concepts(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Comprender")
    count = len(spec.items); cols = 2 if count <= 4 else 3; rows = (count + cols - 1) // cols
    card_w = 11.9 / cols; card_h = 5.25 / max(1, rows)
    size = body_size(spec.items, 18, 13)
    for index, item in enumerate(spec.items):
        col, row = index % cols, index // cols
        x = 0.62 + col * card_w; y = 1.38 + row * card_h
        rounded_card(slide, Inches(x), Inches(y), Inches(card_w - 0.25), Inches(card_h - 0.22))
        add_text(slide, Inches(x + 0.22), Inches(y + 0.18), Inches(0.42), Inches(0.38), f"{index + 1:02}", 15, accent, True)
        add_text(slide, Inches(x + 0.72), Inches(y + 0.18), Inches(card_w - 1.1), Inches(card_h - 0.48), item, size, INK)

def render_focus(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Comprender y transferir")
    concept_count = int(spec.subtitle) if spec.subtitle.isdigit() else 1
    concepts = spec.items[:concept_count]
    examples = spec.items[concept_count:]
    rounded_card(slide, Inches(0.68), Inches(1.45), Inches(5.78), Inches(5.18), WHITE, SOFT)
    add_text(slide, Inches(1.02), Inches(1.78), Inches(2.2), Inches(0.36), "IDEA CLAVE", 13, accent, True)
    add_bullets(slide, Inches(1.02), Inches(2.38), Inches(5.05), Inches(3.75), concepts, body_size(concepts, 23, 17), INK)
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.45), Inches(5.9), Inches(5.18))
    panel.fill.solid(); panel.fill.fore_color.rgb = DARK; panel.line.fill.background()
    add_text(slide, Inches(7.10), Inches(1.78), Inches(2.4), Inches(0.36), "EN PRÁCTICA", 13, accent, True)
    add_bullets(slide, Inches(7.10), Inches(2.38), Inches(5.05), Inches(3.75), examples, body_size(examples, 23, 17), WHITE)

def render_relationship(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Relacionar sin encasillar")
    for index, item in enumerate(spec.items[:4]):
        col, row = index % 2, index // 2
        x, y = 0.68 + col * 6.04, 1.42 + row * 2.18
        rounded_card(slide, Inches(x), Inches(y), Inches(5.72), Inches(1.86))
        add_text(slide, Inches(x + 0.22), Inches(y + 0.18), Inches(0.52), Inches(0.35), f"{index + 1:02}", 14, accent, True)
        add_text(slide, Inches(x + 0.82), Inches(y + 0.16), Inches(4.62), Inches(1.45), item, body_size([item], 16, 13), INK)
    note = spec.items[4] if len(spec.items) > 4 else spec.subtitle
    rounded_card(slide, Inches(0.68), Inches(5.93), Inches(11.95), Inches(0.83), accent, accent)
    add_text(slide, Inches(0.96), Inches(6.12), Inches(11.35), Inches(0.42), note, 16, WHITE, True, PP_ALIGN.CENTER)

def render_debugger(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Mapa del depurador")
    code = rounded_card(slide, Inches(0.66), Inches(1.42), Inches(7.15), Inches(5.22), DARK, DARK)
    add_text(slide, Inches(0.94), Inches(1.65), Inches(5.8), Inches(0.32), "Main.java   ▶  Debug   ↷  Step Over", 13, RGBColor(181, 192, 208), True)
    code_lines = [
        "command = scanner.nextLine();",
        "●  if (command.equals(\"exit\")) {",
        "       running = false;",
        "   }",
        "userName = memory.getUserName();",
    ]
    add_text(slide, Inches(1.12), Inches(2.25), Inches(6.15), Inches(3.5), "\n".join(code_lines), 19, WHITE)
    variables = rounded_card(slide, Inches(8.08), Inches(1.42), Inches(4.56), Inches(5.22), WHITE, SOFT)
    add_text(slide, Inches(8.38), Inches(1.70), Inches(3.95), Inches(0.35), "VARIABLES", 14, accent, True)
    add_text(slide, Inches(8.38), Inches(2.25), Inches(3.82), Inches(1.65), "command  = ____\nrunning  = ____\nuserName = ____", 19, INK, True)
    add_text(slide, Inches(8.38), Inches(4.25), Inches(3.82), Inches(1.7), "\n".join(spec.items[-2:]), 14, MUTED)

def render_pattern(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Decidir desde el problema")
    colors = [RGBColor(56, 94, 142), RGBColor(45, 130, 105), RGBColor(184, 111, 36), RGBColor(122, 83, 153)]
    labels = ["PROBLEMA", "SEMEJANZA", "DIFERENCIA", "DECISIÓN"]
    for index, item in enumerate(spec.items[:4]):
        col, row = index % 2, index // 2
        x, y = 0.66 + col * 6.05, 1.48 + row * 2.55
        rounded_card(slide, Inches(x), Inches(y), Inches(5.72), Inches(2.22), colors[index], colors[index])
        add_text(slide, Inches(x + 0.24), Inches(y + 0.20), Inches(1.35), Inches(0.34), f"{index + 1} · {labels[index]}", 12, WHITE, True)
        add_text(slide, Inches(x + 0.24), Inches(y + 0.68), Inches(5.15), Inches(1.24), item, body_size([item], 17, 13), WHITE, True, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

def render_exception_flow(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Flujo de control del error")
    positions = [(0.55, 1.55), (3.15, 1.55), (5.75, 1.55), (8.35, 1.55), (5.75, 4.30)]
    colors = [RED, AMBER, accent, RGBColor(80, 104, 147), GREEN]
    for index, (item, (x, y)) in enumerate(zip(spec.items, positions)):
        w = 2.35 if index < 4 else 6.85
        rounded_card(slide, Inches(x), Inches(y), Inches(w), Inches(1.65), colors[index], colors[index])
        add_text(slide, Inches(x + 0.18), Inches(y + 0.18), Inches(w - 0.36), Inches(1.26), item, body_size([item], 15, 11), WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if index < 3:
            arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(x + 2.30), Inches(y + 0.59), Inches(0.52), Inches(0.36))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = MUTED; arrow.line.fill.background()
    down = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.DOWN_ARROW, Inches(9.22), Inches(3.37), Inches(0.52), Inches(0.62))
    down.fill.solid(); down.fill.fore_color.rgb = MUTED; down.line.fill.background()

def render_defense(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Rúbrica visual")
    for index, item in enumerate(spec.items):
        col, row = index % 3, index // 3
        x, y = 0.62 + col * 4.13, 1.43 + row * 2.60
        rounded_card(slide, Inches(x), Inches(y), Inches(3.82), Inches(2.22))
        add_text(slide, Inches(x + 0.24), Inches(y + 0.20), Inches(0.55), Inches(0.48), f"{index + 1}", 22, accent, True)
        add_text(slide, Inches(x + 0.88), Inches(y + 0.24), Inches(2.55), Inches(1.55), item, 18, INK, True, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

def render_microdefense(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Modelo de intervención breve")
    widths = [10.9, 10.2, 9.5, 8.8, 8.1]
    for index, item in enumerate(spec.items):
        x = 0.75 + index * 0.36
        y = 1.36 + index * 1.02
        rounded_card(slide, Inches(x), Inches(y), Inches(widths[index]), Inches(0.82), WHITE, accent if index == 0 else SOFT)
        add_text(slide, Inches(x + 0.18), Inches(y + 0.18), Inches(0.44), Inches(0.34), str(index + 1), 15, accent, True, PP_ALIGN.CENTER)
        add_text(slide, Inches(x + 0.72), Inches(y + 0.13), Inches(widths[index] - 0.95), Inches(0.52), item, 16, INK, index == 0)

def render_example(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Observar antes de actuar")
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.68), Inches(1.48), Inches(11.95), Inches(4.95))
    panel.fill.solid(); panel.fill.fore_color.rgb = DARK; panel.line.fill.background()
    add_text(slide, Inches(1.02), Inches(1.78), Inches(2.25), Inches(0.45), "DEMOSTRACIÓN", 13, accent, True)
    add_bullets(slide, Inches(1.02), Inches(2.42), Inches(10.9), Inches(3.55), spec.items, body_size(spec.items, 22, 15), WHITE)

def render_activity(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Aplicar y decidir")
    size = body_size(spec.items, 20, 14)
    start_step = int(spec.subtitle) if spec.subtitle.isdigit() else 1
    for index, item in enumerate(spec.items):
        y = 1.43 + index * (5.4 / max(1, len(spec.items)))
        block_h = 5.15 / max(1, len(spec.items))
        rounded_card(slide, Inches(0.72), Inches(y), Inches(11.85), Inches(block_h - 0.14))
        add_text(slide, Inches(0.93), Inches(y + 0.16), Inches(0.75), Inches(0.42), f"PASO {start_step + index}", 12, accent, True)
        add_text(slide, Inches(1.82), Inches(y + 0.12), Inches(10.35), Inches(block_h - 0.38), item, size, INK)

def render_evidence(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Comprobar")
    size = body_size(spec.items, 19, 14); y = 1.42; step = 5.45 / max(1, len(spec.items))
    for item in spec.items:
        square = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.78), Inches(y), Inches(0.46), Inches(0.46))
        square.fill.solid(); square.fill.fore_color.rgb = WHITE; square.line.color.rgb = accent
        add_text(slide, Inches(1.52), Inches(y - 0.02), Inches(10.75), Inches(step - 0.04), item, size, INK)
        y += step

def render_safety(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Límites que protegen el aprendizaje")
    colors = [GREEN, AMBER, RED, accent]
    cols = 2; rows = (len(spec.items) + 1) // 2; card_h = 5.2 / max(1, rows)
    size = body_size(spec.items, 19, 14)
    for index, item in enumerate(spec.items):
        col, row = index % cols, index // cols
        x = 0.68 + col * 6.02; y = 1.42 + row * card_h
        rounded_card(slide, Inches(x), Inches(y), Inches(5.72), Inches(card_h - 0.22))
        marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 0.24), Inches(y + 0.25), Inches(0.26), Inches(card_h - 0.72))
        marker.fill.solid(); marker.fill.fore_color.rgb = colors[index % len(colors)]; marker.line.fill.background()
        add_text(slide, Inches(x + 0.72), Inches(y + 0.25), Inches(4.62), Inches(card_h - 0.65), item, size, INK)

def render_closure(prs: Presentation, session: Session, spec: SlideSpec) -> None:
    slide, accent = base_slide(prs, session, spec.title, "Explicar y defender")
    if spec.items:
        question = spec.items[-1]
        supporting = spec.items[:-1]
    else:
        question = session.close_question
        supporting = []
    box = rounded_card(slide, Inches(0.68), Inches(1.42), Inches(11.95), Inches(2.0), accent, accent)
    add_text(slide, Inches(1.02), Inches(1.74), Inches(11.2), Inches(0.28), "PREGUNTA O DECISIÓN DE SALIDA", 13, WHITE, True)
    add_text(slide, Inches(1.02), Inches(2.16), Inches(11.0), Inches(0.9), question, body_size([question], 25, 17), WHITE, True)
    if supporting:
        add_bullets(slide, Inches(0.82), Inches(3.92), Inches(11.6), Inches(2.35), supporting, body_size(supporting, 18, 14))
    else:
        add_text(slide, Inches(0.82), Inches(4.25), Inches(11.6), Inches(1.4), "La sesión termina cuando la evidencia existe, está comprobada y puede defenderse.", 23, GREEN, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

RENDERERS = {
    "title": render_title,
    "outcome": render_outcome,
    "timeline": render_timeline,
    "concepts": render_concepts,
    "focus": render_focus,
    "relationship": render_relationship,
    "debugger": render_debugger,
    "pattern": render_pattern,
    "exception_flow": render_exception_flow,
    "defense": render_defense,
    "microdefense": render_microdefense,
    "example": render_example,
    "activity": render_activity,
    "evidence": render_evidence,
    "safety": render_safety,
    "closure": render_closure,
}

def speaker_notes(session: Session, spec: SlideSpec, index: int, total: int) -> str:
    """Genera un guion docente trazable desde la misma fuente que la diapositiva."""
    lines = [
        f"Objetivo docente: {session.objective}",
        f"Diapositiva {index} de {total}: {spec.title}",
    ]
    if session.moment:
        lines.append(f"Fase HEXA: {session.moment}")

    if spec.kind == "title":
        lines.append(f"Presenta el producto y anticipa la evidencia: {session.evidence}")
        if session.folder == "h1":
            lines.append("Secuencia docente:")
            lines.extend(f"- {block.time}: {block.action}" for block in session.timeline)
            if session.key_concepts:
                lines.append("Claves docentes:")
                lines.extend(f"- {item}" for item in session.key_concepts)
    elif spec.kind == "timeline":
        lines.append("Conduce estos tiempos y evita ampliar la explicación si reduce la práctica:")
        lines.extend(f"- {block.time}: {block.action}" for block in session.timeline)
    elif spec.kind in {"concepts", "focus", "relationship", "debugger", "pattern", "exception_flow"}:
        lines.append("Explica con predicción y ejemplo mínimo; comprueba comprensión antes de continuar.")
        lines.extend(f"- {item}" for item in spec.items)
    elif spec.kind == "example":
        lines.append("Modela el ejemplo sin resolver la actividad completa; pide una predicción antes de ejecutar.")
        lines.extend(f"- {item}" for item in spec.items)
    elif spec.kind == "activity":
        lines.append("Entrega la consigna, observa decisiones y pide una prueba reproducible.")
        lines.extend(f"- {item}" for item in spec.items)
    elif spec.kind == "evidence":
        lines.append(f"Comprueba la evidencia en su fuente canónica: {session.evidence}")
        lines.extend(f"- {item}" for item in spec.items)
    elif spec.kind == "safety":
        lines.append("Recuerda estos límites antes de que el alumnado comparta o publique:")
        lines.extend(f"- {item}" for item in spec.items)
    elif spec.kind in {"closure", "defense", "microdefense"}:
        lines.append("Cierra con explicación individual, prueba observable y siguiente paso concreto.")
        lines.extend(f"- {item}" for item in spec.items)
    else:
        lines.extend(f"- {item}" for item in spec.items)

    lines.append(f"Fuente docente: {session.teacher_source.as_posix()}")
    return "\n".join(lines)

def add_speaker_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text

def build_presentation(session: Session, target: Path) -> int:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    plan = plan_slides(session)
    for index, spec in enumerate(plan, 1):
        RENDERERS[spec.kind](prs, session, spec)
        add_speaker_notes(prs.slides[-1], speaker_notes(session, spec, index, len(plan)))
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(target)
    return len(plan)

def target_for(session: Session, output_root: Path) -> Path:
    stem = session.teacher_source.name.replace("-docente.md", "-presentacion.pptx")
    return output_root / session.folder / stem

def collect_sessions(selected: list[str] | None = None) -> list[Session]:
    wanted = {value.upper().removeprefix("S") for value in selected or []}
    sessions: list[Session] = []
    for teacher_path in sorted(TEACHER_ROOT.rglob("S*-docente.md")):
        match = re.search(r"S(\d{3})", teacher_path.name)
        if not match or (wanted and match.group(1) not in wanted):
            continue
        sessions.append(parse_session(teacher_path, student_for_teacher(teacher_path)))
    sessions.sort(key=lambda item: int(item.number))
    return sessions

def generate_index(generated: list[tuple[Session, Path, int]], output_root: Path) -> None:
    lines = [
        "# Presentaciones por sesión - MiniJarvis",
        "",
        "Presentaciones adaptativas generadas principalmente desde las guías docentes.",
        "",
        "| Sesión | Hito | Tema | Diapositivas | Presentación |",
        "|---:|---|---|---:|---|",
    ]
    for session, target, count in generated:
        relative = target.relative_to(output_root).as_posix()
        lines.append(f"| {session.number} | {session.hito} | {session.topic} | {count} | `{relative}` |")
    (output_root / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

def atomic_publish(staging: Path, output: Path, backup_root: Path) -> Path | None:
    """Publica staging conservando el directorio anterior completo en backup_root."""
    backup_root.mkdir(parents=True, exist_ok=True)
    backup: Path | None = None
    if output.exists():
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup = backup_root / f"{output.name}-{stamp}"
        counter = 1
        while backup.exists():
            backup = backup_root / f"{output.name}-{stamp}-{counter}"
            counter += 1
        shutil.move(str(output), str(backup))
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staging), str(output))
    except Exception:
        if backup is not None and backup.exists() and not output.exists():
            shutil.move(str(backup), str(output))
        raise
    return backup

def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Valida fuentes y planes sin escribir PPTX.")
    parser.add_argument("--dry-run", action="store_true", help="Muestra el plan de diapositivas sin generar archivos.")
    parser.add_argument("--session", dest="sessions", action="append", help="Limita la operación a una sesión, por ejemplo S203. Repetible.")
    parser.add_argument("--output-dir", type=Path, help="Directorio alternativo. Debe estar vacío o no existir.")
    parser.add_argument("--backup-dir", type=Path, help="Directorio temporal donde conservar la salida anterior.")
    return parser

def validate_cli_safety(args: argparse.Namespace) -> None:
    if args.sessions and args.output_dir is None and not args.check and not args.dry_run:
        raise SystemExit("La generación parcial con --session requiere --output-dir para no sustituir la colección completa.")

def print_validation(sessions: list[Session]) -> bool:
    valid = True
    for session in sessions:
        report = validate_session(session)
        status = "PASS" if not report.errors else "FAIL"
        print(f"S{session.number}: {status} slides={len(plan_slides(session))}")
        for error in report.errors:
            print(f"  ERROR: {error}")
        for warning in report.warnings:
            print(f"  WARNING: {warning}")
        valid = valid and not report.errors
    return valid

def generate_to(sessions: list[Session], output_root: Path) -> list[tuple[Session, Path, int]]:
    generated: list[tuple[Session, Path, int]] = []
    for session in sessions:
        target = target_for(session, output_root)
        count = build_presentation(session, target)
        generated.append((session, target, count))
    generate_index(generated, output_root)
    return generated

def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    validate_cli_safety(args)
    sessions = collect_sessions(args.sessions)
    if not sessions:
        raise SystemExit("No se encontraron sesiones para la selección indicada.")

    valid = print_validation(sessions)
    if args.check and not valid:
        return 1
    if args.check and not args.dry_run:
        return 0

    if args.dry_run:
        for session in sessions:
            print(f"S{session.number}: " + " -> ".join(slide.kind for slide in plan_slides(session)))
        return 0 if valid else 1

    if not valid:
        print("Generación cancelada: hay errores de validación.")
        return 1

    if args.output_dir is not None:
        output = args.output_dir.resolve()
        if output.exists() and any(output.iterdir()):
            raise SystemExit(f"El directorio alternativo no está vacío: {output}")
        output.mkdir(parents=True, exist_ok=True)
        generated = generate_to(sessions, output)
        print(f"Presentaciones generadas: {len(generated)} en {output}")
        return 0

    staging = Path(tempfile.mkdtemp(prefix="minijarvis-presentaciones-staging-")) / OUTPUT_ROOT.name
    backup_root = args.backup_dir or Path(tempfile.mkdtemp(prefix="minijarvis-presentaciones-backup-"))
    generated = generate_to(sessions, staging)
    backup = atomic_publish(staging, OUTPUT_ROOT, backup_root)
    print(f"Presentaciones generadas: {len(generated)}")
    if backup:
        print(f"Salida anterior conservada en: {backup}")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
