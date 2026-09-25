#!/usr/bin/env python3
"""Genera las cinco evidencias digitales de Laura para cada hito."""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
EXAMPLES = ROOT / "03-EJEMPLOS-LAURA-PRIVADOS"
MASTERS = ROOT / "02-PROFESORADO" / "05-ECOSISTEMA-DIGITAL" / "PLANTILLAS-MAESTRAS"

HITOS = {
    "h0-torre-papel": {
        "code": "H0",
        "title": "Torre de papel y Scrum",
        "objective": "Organizar un sprint, construir y comprobar una torre de papel.",
        "work": "Preparé el backlog, asumí una responsabilidad, registré bloqueos y participé en la revisión y retrospectiva.",
        "test": "La torre permaneció en pie durante 10 segundos y medimos su altura.",
        "result": "El equipo completó el sprint y acordó una mejora para MiniJarvis.",
        "problem": "La base inicial era demasiado estrecha.",
        "decision": "Ensanchamos la base antes de aumentar la altura.",
        "increment": "Torre estable, backlog actualizado y retrospectiva del equipo.",
        "next": "Dividir mejor las tareas del primer hito de programación.",
        "version": "No aplica GitHub en H0",
    },
    "h1-primer-asistente": {
        "code": "H1",
        "title": "Primer asistente por consola",
        "objective": "Crear y ejecutar la primera versión básica de MiniJarvis.",
        "work": "Programé la entrada del nombre, las constantes y los mensajes de consola; después comprobé la ejecución.",
        "test": "Ejecuté el programa con un nombre normal y con una entrada vacía.",
        "result": "MiniJarvis saluda, solicita el nombre y termina sin errores.",
        "problem": "Confundí texto literal con el nombre de una variable.",
        "decision": "Mantuve H1 sin menú ni bucles para poder defender cada línea.",
        "increment": "Primer MiniJarvis ejecutable con README y evidencia de ejecución.",
        "next": "Añadir un menú controlado en H2.",
        "version": "h1-entrega",
    },
    "h2-decisiones-depuracion": {
        "code": "H2",
        "title": "Decisiones y depuración",
        "objective": "Añadir menú, repetición, decisiones y salida controlada.",
        "work": "Implementé comandos, preparé pruebas manuales y documenté una depuración con breakpoint.",
        "test": "Probé ayuda, saludo, estado, comando desconocido y salir.",
        "result": "El menú se repite hasta salir y responde de forma controlada.",
        "problem": "El programa comparaba textos con el operador incorrecto.",
        "decision": "Usamos equalsIgnoreCase para comparar comandos.",
        "increment": "MiniJarvis interactivo con pruebas e incidencia documentada.",
        "next": "Añadir memoria temporal mediante una colección.",
        "version": "h2-entrega",
    },
    "h3-memoria-colecciones": {
        "code": "H3",
        "title": "Memoria y colecciones",
        "objective": "Guardar y consultar recuerdos durante una ejecución.",
        "work": "Implementé los comandos recuerda y memoria y documenté por qué usamos una colección.",
        "test": "Probé memoria vacía, un recuerdo, varios recuerdos y entradas repetidas.",
        "result": "Los recuerdos permanecen disponibles hasta cerrar el programa.",
        "problem": "La memoria vacía no mostraba un mensaje claro.",
        "decision": "Usamos una lista y añadimos un caso específico para memoria vacía.",
        "increment": "MiniJarvis con memoria temporal y pruebas de casos límite.",
        "next": "Separar responsabilidades mediante clases.",
        "version": "h3-entrega",
    },
    "h4-agente-orientado-objetos": {
        "code": "H4",
        "title": "Agente orientado a objetos",
        "objective": "Reorganizar MiniJarvis con clases y responsabilidades claras.",
        "work": "Separé Main, Agent y Memory y relacioné el código con los diagramas.",
        "test": "Repetí las pruebas de H2 y H3 después de reorganizar el código.",
        "result": "El comportamiento anterior se conserva con un diseño más claro.",
        "problem": "Agent acumulaba responsabilidades de entrada, memoria y comandos.",
        "decision": "Main inicia la aplicación, Agent coordina y Memory conserva recuerdos.",
        "increment": "Código orientado a objetos con diagramas coherentes.",
        "next": "Facilitar la incorporación de herramientas nuevas.",
        "version": "h4-entrega",
    },
    "h5-extensible-clean-code-patrones": {
        "code": "H5",
        "title": "Extensibilidad, código limpio y patrones",
        "objective": "Añadir herramientas con bajo impacto y refactorizar sin romper el programa.",
        "work": "Implementé la interfaz Tool, una herramienta nueva y una revisión de código.",
        "test": "Ejecuté pruebas de regresión antes y después de la refactorización.",
        "result": "Las herramientas comparten un contrato y el menú sigue funcionando.",
        "problem": "Cada comando nuevo obligaba a modificar un bloque grande de condiciones.",
        "decision": "Aplicamos un Command simplificado mediante la interfaz Tool.",
        "increment": "MiniJarvis extensible con refactorización y revisión Git.",
        "next": "Guardar memoria e historial entre ejecuciones.",
        "version": "h5-entrega",
    },
    "h6-persistente-trazable": {
        "code": "H6",
        "title": "Persistencia y trazabilidad",
        "objective": "Conservar memoria e historial entre ejecuciones.",
        "work": "Implementé lectura y escritura de ficheros, tratamiento de errores y un historial sin datos sensibles.",
        "test": "Guardé un recuerdo, cerré el programa y comprobé que aparecía al volver a ejecutarlo.",
        "result": "La memoria persiste y los errores de fichero se controlan.",
        "problem": "La primera versión dependía de una ruta absoluta.",
        "decision": "Usamos rutas relativas y datos ficticios.",
        "increment": "MiniJarvis persistente con pruebas, seguridad e historial.",
        "next": "Validar una ayuda de IA simulada o autorizada.",
        "version": "h6-entrega",
    },
    "h7-integracion-ia-responsable": {
        "code": "H7",
        "title": "Integración responsable de IA",
        "objective": "Incorporar una ayuda de IA simulada con límites y validación humana.",
        "work": "Definí el caso de uso, filtré entradas inseguras y registré prompts y validaciones.",
        "test": "Probé una petición válida, una entrada con datos sensibles y una respuesta incorrecta.",
        "result": "La simulación rechaza riesgos y exige validar las respuestas.",
        "problem": "Una respuesta simulada parecía válida pero contenía una afirmación falsa.",
        "decision": "Toda respuesta se muestra como propuesta y requiere validación humana.",
        "increment": "Ayuda de IA simulada, segura, registrada y defendible.",
        "next": "Preparar la demostración y defensa final.",
        "version": "h7-entrega",
    },
    "hf-presentacion-final-recuperacion-mejora": {
        "code": "HF",
        "title": "Presentación final, recuperación y mejora",
        "objective": "Seleccionar evidencias y demostrar la evolución de MiniJarvis.",
        "work": "Organicé el portfolio final, preparé la demostración y ensayé la defensa individual.",
        "test": "Seguí el guion de instalación y ejecución desde un entorno limpio.",
        "result": "La demostración es reproducible y las evidencias permiten explicar el progreso.",
        "problem": "Una evidencia antigua no identificaba la versión exacta del código.",
        "decision": "Enlazamos cada evidencia con su commit o tag estable.",
        "increment": "Portfolio, demostración, defensa y plan de mejora final.",
        "next": "Aplicar el plan de mejora después de recibir retroalimentación.",
        "version": "hf-final",
    },
}

def set_row(ws, row: int, values: tuple[object, ...]) -> None:
    for column, value in enumerate(values, start=1):
        ws.cell(row=row, column=column, value=value)

def create_diary(target: Path, data: dict[str, str]) -> None:
    source = MASTERS / "01-Diario-individual-MiniJarvis.xlsx"
    shutil.copy2(source, target)
    workbook = load_workbook(target)
    sheet = workbook["DIARIO"]
    set_row(
        sheet,
        2,
        (
            "Fecha de ejemplo",
            data["code"],
            data["code"],
            data["objective"],
            data["work"],
            data["test"],
            "URL_RESTRINGIDA_EJEMPLO",
            data["result"],
            data["problem"],
            "Explicación y revisión de claridad",
            "Contrasté la respuesta con el código, la ejecución y las pruebas",
            data["next"],
        ),
    )
    workbook.save(target)

def create_scrum(target: Path, data: dict[str, str]) -> None:
    source = MASTERS / "02-Scrum-equipo-MiniJarvis.xlsx"
    shutil.copy2(source, target)
    workbook = load_workbook(target)
    code = data["code"]
    people = (
        ("Equipo Ada", "Laura", "Documentación y defensa", "Avisar de bloqueos", "Rotar en el siguiente hito"),
        ("Equipo Ada", "Alex", "Responsable técnico", "Revisar antes de integrar", "Rotar en el siguiente hito"),
        ("Equipo Ada", "Noor", "Backlog y pruebas", "Cada tarea deja evidencia", "Rotar en el siguiente hito"),
    )
    for row, values in enumerate(people, start=2):
        set_row(workbook["EQUIPO_Y_ACUERDOS"], row, values)

    tasks = (
        (f"{code}-01", code, "Construir el incremento", "Funciona y se puede explicar", "Alta", "Alex", "Hecho", "URL_RESTRINGIDA_EJEMPLO"),
        (f"{code}-02", code, "Probar casos principales", "Resultados registrados", "Alta", "Noor", "Hecho", "URL_RESTRINGIDA_EJEMPLO"),
        (f"{code}-03", code, "Seleccionar evidencias", "Enlaces comprobados", "Media", "Laura", "Hecho", "URL_RESTRINGIDA_EJEMPLO"),
    )
    for row, values in enumerate(tasks, start=2):
        set_row(workbook["PRODUCT_BACKLOG"], row, values)
        set_row(
            workbook["SPRINT_ACTUAL"],
            row,
            (values[0], values[2], values[5], "Inicio de ejemplo", "Fin de ejemplo", values[6], "", values[3]),
        )

    set_row(workbook["DECISIONES"], 2, ("Fecha de ejemplo", code, data["decision"], "Otras opciones discutidas", "Reduce riesgo y facilita la defensa", "Todo el equipo", "URL_RESTRINGIDA_EJEMPLO"))
    set_row(workbook["BLOQUEOS"], 2, ("Fecha de ejemplo", code, data["problem"], "Retrasaba la prueba", "Reducir el problema y probar una alternativa", "Laura", "Resuelto"))
    set_row(workbook["REVIEW"], 2, ("Fecha de ejemplo", code, data["increment"], data["test"], data["result"], "Retroalimentación de ejemplo", data["next"], "URL_RESTRINGIDA_EJEMPLO"))
    set_row(workbook["RETROSPECTIVAS"], 2, ("Fecha de ejemplo", code, "Tareas pequeñas y pruebas visibles", "Avisar antes de que un bloqueo crezca", data["next"], "Equipo Ada", "Siguiente hito"))
    set_row(workbook["REGISTRO_IA_EQUIPO"], 2, ("Fecha de ejemplo", code, "IA autorizada", "Revisar claridad", "Revisa esta explicación sin cambiar el contenido técnico", "Solo sugerencias de redacción", "Comparación con código y pruebas", "Laura", "URL_RESTRINGIDA_EJEMPLO"))

    links = [
        (code, "Site personal", f"Página {code} de Laura", "URL_RESTRINGIDA_EJEMPLO", "Sí", code, "PDF si se solicita"),
        (code, "Site de equipo", f"Incremento {code} del Equipo Ada", "URL_RESTRINGIDA_EJEMPLO", "Sí", code, "PDF si se solicita"),
        (code, "Diario individual", f"Filas {code}", "URL_RESTRINGIDA_EJEMPLO", "Sí", "Fecha de ejemplo", "XLSX"),
        (code, "Sheet Scrum", f"Filas {code}", "URL_RESTRINGIDA_EJEMPLO", "Sí", "Fecha de ejemplo", "XLSX"),
        (code, "GitHub" if code != "H0" else "Drive", data["version"], "URL_RESTRINGIDA_EJEMPLO", "Sí", data["version"], "No aplica"),
    ]
    for row, values in enumerate(links, start=2):
        set_row(workbook["ENLACES_EVIDENCIAS"], row, values)
    workbook.save(target)

def site_personal(data: dict[str, str]) -> str:
    return f"""# Ejemplo de Site personal de Laura - {data['code']}

## Página: {data['code']} - {data['title']}

### Qué hice yo

{data['work']}

### Qué aprendí

{data['objective']}

### Problema encontrado

{data['problem']}

### Prueba realizada

{data['test']}

### Evidencia seleccionada

Enlace profundo de ejemplo: `URL_RESTRINGIDA_EJEMPLO`.

He elegido esta evidencia porque permite comprobar el resultado y explicar mi aportación.

### Qué puedo defender

Puedo explicar la decisión principal, mostrar la prueba y localizar mi aportación.

### Uso de IA y validación

Usé IA solo para revisar una explicación. Validé las sugerencias comparándolas con el trabajo y las pruebas reales.

### Siguiente mejora

{data['next']}
"""

def site_team(data: dict[str, str]) -> str:
    github = "No se usa GitHub en H0." if data["code"] == "H0" else f"Versión evaluada: `{data['version']}`."
    return f"""# Ejemplo de Site del Equipo Ada - {data['code']}

## Reto e incremento

Reto: {data['objective']}

Incremento: {data['increment']}

## Decisión principal

{data['decision']}

## Pruebas realizadas

{data['test']}

Resultado: {data['result']}

## Evidencias seleccionadas

- Sheet Scrum: `URL_RESTRINGIDA_EJEMPLO`.
- Evidencia del incremento: `URL_RESTRINGIDA_EJEMPLO`.
- {github}

## Revisión

El equipo mostró el incremento, reprodujo la prueba principal y anotó la retroalimentación recibida.

## Retrospectiva y siguiente mejora

Detectamos este problema: {data['problem']}

Siguiente acción: {data['next']}
"""

def moodle_delivery(data: dict[str, str]) -> str:
    github = "No aplica" if data["code"] == "H0" else data["version"]
    return f"""# Ejemplo de entrega de enlaces en Moodle - {data['code']}

## Identificación

- Hito: {data['code']} - {data['title']}.
- Equipo: Equipo Ada.
- Alumna: Laura.
- Fecha: fecha de ejemplo.

## Enlaces entregados

| Evidencia | URL concreta | Permiso comprobado | Identificador estable |
|---|---|---|---|
| Site personal | `URL_RESTRINGIDA_EJEMPLO` | Sí | Página {data['code']} |
| Site de equipo | `URL_RESTRINGIDA_EJEMPLO` | Sí | Página {data['code']} |
| Diario individual | `URL_RESTRINGIDA_EJEMPLO` | Sí | Filas {data['code']} |
| Sheet Scrum | `URL_RESTRINGIDA_EJEMPLO` | Sí | Filas {data['code']} |
| Drive | `URL_RESTRINGIDA_EJEMPLO` | Sí | Carpeta {data['code']} |
| GitHub, si procede | `URL_RESTRINGIDA_EJEMPLO` | Sí | {github} |

## Evidencia cerrada

- Archivo: `{data['code']}-equipo-ada-evidencias.pdf`.
- Contenido: selección de evidencias y resultados del hito.

## Declaración

Hemos comprobado los permisos con una cuenta distinta de la propietaria. Los enlaces son ejemplos restringidos, no contienen datos personales ni secretos y podemos explicar lo entregado.
"""

def update_readme(path: Path, data: dict[str, str]) -> None:
    marker = "## Evidencias digitales correspondientes a la entrega"
    text = path.read_text(encoding="utf-8") if path.exists() else f"# {data['code']} - {data['title']}\n"
    if marker in text:
        text = text.split(marker, 1)[0].rstrip() + "\n"
    text += f"""

{marker}

La carpeta `evidencias-digitales` muestra cómo se presenta este hito en los cinco documentos comunes del curso:

1. `01-Diario-individual-MiniJarvis.xlsx`.
2. `02-Scrum-equipo-MiniJarvis.xlsx`.
3. `03-Site-personal-estructura.md`.
4. `04-Site-equipo-estructura.md`.
5. `05-Entrega-enlaces-Moodle.md`.

Los archivos de código y la carpeta `docs` contienen las evidencias técnicas originales. Los cinco documentos anteriores las seleccionan, explican y entregan; no las sustituyen.
"""
    path.write_text(text, encoding="utf-8")

def main() -> None:
    for folder, data in HITOS.items():
        hito = EXAMPLES / folder
        if not hito.exists():
            raise RuntimeError(f"No existe el hito: {hito}")
        evidence = hito / "evidencias-digitales"
        evidence.mkdir(exist_ok=True)
        create_diary(evidence / "01-Diario-individual-MiniJarvis.xlsx", data)
        create_scrum(evidence / "02-Scrum-equipo-MiniJarvis.xlsx", data)
        (evidence / "03-Site-personal-estructura.md").write_text(site_personal(data), encoding="utf-8")
        (evidence / "04-Site-equipo-estructura.md").write_text(site_team(data), encoding="utf-8")
        (evidence / "05-Entrega-enlaces-Moodle.md").write_text(moodle_delivery(data), encoding="utf-8")
        update_readme(hito / "README.md", data)

if __name__ == "__main__":
    main()
