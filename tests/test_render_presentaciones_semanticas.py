from __future__ import annotations

import json
import re
from collections import Counter

from pptx import Presentation
from pptx.util import Inches
import pytest

from test_generar_presentaciones_sesiones import load_module, source_pair


@pytest.fixture(scope="module")
def generator():
    return load_module()


def metadata(slide):
    nodes = slide._element.xpath(".//*[local-name()='renderMetadata']")
    assert len(nodes) == 1, "Falta metadata semántica persistida"
    return json.loads(nodes[0].text)


def visible_text(slide):
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


@pytest.mark.parametrize("number", [str(number) for number in range(206, 216)])
def test_h1_pptx_projects_literal_student_mode(generator, tmp_path, number):
    teacher, student = source_pair(number)
    literal = re.search(
        r"\*\*Modalidad:\*\* (.+?)\.", student.read_text(encoding="utf-8")
    ).group(1)
    target = tmp_path / f"S{number}.pptx"

    generator.build_presentation(generator.parse_session(teacher, student), target)

    deck_text = "\n".join(visible_text(slide) for slide in Presentation(target).slides)
    assert literal in deck_text


def test_semantic_title_is_not_overloaded_and_has_real_notes(generator, tmp_path):
    session = generator.parse_session(*source_pair("206"))
    target = tmp_path / "S206.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    first = deck.slides[0]
    assert first.has_notes_slide, "Faltan notas reales del presentador"
    data = metadata(first)
    assert data["layout"] == "title"
    assert data["unit_ids"] and data["source_refs"]
    assert data["timing"]["seconds"] >= 0
    text = visible_text(first)
    assert "206" in text and session.topic in text
    assert "INDIVIDUAL → EQUIPO" in text
    assert "La sesión permite seguir este recorrido" not in text
    assert "Apertura" in first.notes_slide.notes_text_frame.text


def test_scanner_is_editable_complete_and_every_frame_has_trace_and_time(generator, tmp_path):
    session = generator.parse_session(*source_pair("212"))
    plan = generator.plan_slides(session)
    target = tmp_path / "S212.pptx"
    count = generator.build_presentation(session, target)
    deck = Presentation(target)
    assert len(deck.slides) == count
    assert abs(deck.slide_width / deck.slide_height - 16 / 9) < 0.001
    data = [metadata(slide) for slide in deck.slides]
    assert sum(record["timing"]["seconds"] for record in data) == 45 * 60
    # seconds is accounting ownership, not a compulsory dwell time per screen.
    assert all(record["timing"]["seconds"] >= 0 for record in data)
    assert all(record["timing"]["shared_intervals"] for record in data)
    assert all(interval["seconds"] > 0 for record in data for interval in record["timing"]["shared_intervals"])
    assert all(record["unit_ids"] and record["source_refs"] for record in data)
    shapes = [shape for slide in deck.slides for shape in slide.shapes]
    assert all(shape.shape_type != 13 for shape in shapes), "El código no puede ser imagen"
    program = next(text for spec in plan for text in spec.visible_content
                   if text.startswith("```java\nimport java.util.Scanner"))
    body = program.split("\n", 1)[1].rsplit("\n", 1)[0]
    assert any(shape.has_text_frame and shape.text == body for shape in shapes)
    assert all(slide.has_notes_slide for slide in deck.slides)
    secondary = [(slide, record) for slide, record in zip(deck.slides, data) if record["recognition"]]
    assert secondary
    assert any("Boolean.parseBoolean" in visible_text(slide) and "Para reconocer" in visible_text(slide)
               for slide, _ in secondary)


def test_parallel_columns_are_real_and_bank_stays_in_notes(generator, tmp_path):
    session = generator.parse_session(*source_pair("215"))
    target = tmp_path / "S215.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    parallel = [slide for slide in deck.slides if metadata(slide)["layout"] == "parallel"]
    assert len(parallel) == 1
    slide = parallel[0]
    text = visible_text(slide)
    assert all(word in text for word in ("DOCENTE", "PAREJAS", "EQUIPO", "Defensa", "Ensayo", "Review", "Retrospectiva"))
    lanes = [shape for shape in slide.shapes if shape.name.startswith("Carril ")]
    assert len({shape.left for shape in lanes}) == 3
    assert len({shape.top for shape in lanes}) == 1
    record = metadata(slide)
    assert record["relations"]["concurrency"]["intervals"] == [
        {"time": block.time, "teacher": block.action, "details": block.details} for block in session.timeline]
    notes = "\n".join(s.notes_slide.notes_text_frame.text for s in deck.slides)
    projected = "\n".join(visible_text(s) for s in deck.slides)
    for spec in generator.plan_slides(session):
        for unit in spec.pedagogical_units:
            if unit.function == "question_bank":
                assert all(text in notes for text in unit.presenter_content)
    assert "Banco de preguntas" in notes
    assert "registro nominativo" not in projected.lower()
    assert sum(metadata(s)["timing"]["seconds"] for s in deck.slides) == 2700


def test_classification_keeps_all_proposals_and_empty_destinations_together(generator, tmp_path):
    session = generator.parse_session(*source_pair("206"))
    target = tmp_path / "S206.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    candidates = [slide for slide in deck.slides if metadata(slide)["layout"] == "classification"]
    assert len(candidates) == 1
    text = visible_text(candidates[0])
    for phrase in ("saludar", "pedir un nombre ficticio", "hacer un menú", "recordar conversaciones",
                   "guardar en fichero", "usar una API de IA", "entra en H1", "más adelante", "fuera del reto"):
        assert phrase in text
    assert "sin asignar" in text
    assert metadata(deck.slides[-1])["layout"] == "closure"


def test_activity_lists_are_editable_cards_not_a_generic_bullet_box(generator, tmp_path):
    session = generator.parse_session(*source_pair("213"))
    target = tmp_path / "S213.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    cards = [shape for slide in deck.slides for shape in slide.shapes if shape.name.startswith("Acción o criterio ")]
    assert len(cards) >= 6
    assert len({shape.left for shape in cards}) >= 2
    assert any("INDIVIDUAL" in visible_text(s) and "QUÉ COMPROBAR" in visible_text(s) for s in deck.slides)


def normalized(text):
    text = re.sub(r"(?m)^\s*(?:[-*]\s+|\d+\.\s+|>\s*)", "", text)
    text = text.replace("**", "").replace("`", "")
    return re.sub(r"\s+", " ", text).strip()


@pytest.mark.parametrize("number", ["206", "208", "209", "210", "211", "212", "213", "214", "215"])
def test_artifact_conserves_plan_channels_code_and_every_unit(generator, tmp_path, number):
    import render_presentaciones_semanticas as renderer

    session = generator.parse_session(*source_pair(number))
    plan = generator.plan_slides(session)
    target = tmp_path / f"S{number}.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    records = [metadata(slide) for slide in deck.slides]
    projected = normalized("\n".join(visible_text(slide) for slide in deck.slides))
    notes = normalized("\n".join(slide.notes_slide.notes_text_frame.text for slide in deck.slides))
    presentation_normalizations = {
        original: corrected for record in records
        for original, corrected in record["relations"].get("presentation_normalizations", {}).items()
    }
    unit_ids = {key for record in records for key in record["unit_ids"]}
    expected_ids = {unit.unit_id for spec in plan for unit in spec.pedagogical_units}
    assert unit_ids == expected_ids
    java = []
    for spec_index, spec in enumerate(plan):
        for original in spec.visible_content:
            if (re.fullmatch(r"\s*(?:Pregunta(?: principal| complementaria)?|Cierra en voz alta|Cuando proceda|Di en voz alta|Pregunta antes de probar):\s*", original, re.I)
                    or "\n" not in original.strip() and original.rstrip().endswith(":")):
                assert any(source["sha256"] == __import__("hashlib").sha256(original.encode()).hexdigest()
                           for record in records for source in record["notes_sources"])
                continue
            fenced = renderer.fenced_body(original)
            body = fenced[1] if fenced else original
            if normalized(body) not in projected and normalized(original) not in notes:
                # Diagramas, campos y tarjetas conservan los componentes aunque
                # títulos/pies separen las vistas; Java se exige íntegro abajo.
                fragments = body.split("|") if "|" in body and "||" not in body else body.splitlines()
                for fragment in fragments:
                    if fragment.strip():
                        transformed = any(normalized(fragment) in normalized(original)
                                          and normalized(corrected) in projected
                                          for original, corrected in presentation_normalizations.items())
                        assert normalized(fragment) in projected or normalized(fragment) in notes or transformed, (number, spec_index, fragment)
            if fenced and fenced[0] == "java":
                java.append(body)
        for original in spec.presenter_content:
            if (re.fullmatch(r"\s*(?:Pregunta(?: principal| complementaria)?|Cierra en voz alta|Cuando proceda|Di en voz alta|Pregunta antes de probar):\s*", original, re.I)
                    or "\n" not in original.strip() and original.rstrip().endswith(":")):
                # Empty source labels remain traceable in notes_sources metadata;
                # the new classroom notes contract deliberately does not print them.
                assert any(source["sha256"] == __import__("hashlib").sha256(original.encode()).hexdigest()
                           for record in records for source in record["notes_sources"])
                continue
            fenced = renderer.fenced_body(original)
            body = fenced[1] if fenced else original
            if normalized(body) not in notes and normalized(body) not in projected:
                for fragment in body.splitlines():
                    if fragment.strip():
                        assert normalized(fragment) in notes or normalized(fragment) in projected, (
                            number, spec_index, fragment,
                        )
            if not any(original == text or original in text for text in spec.visible_content):
                assert normalized(original) not in projected, (number, "presenter-only filtrado a proyección", original)
    actual_code = [body for record in records for body in record["code_blocks"]]
    assert Counter(actual_code) == Counter(java)
    for body in java:
        matching = [shape for slide in deck.slides for shape in slide.shapes if shape.has_text_frame and shape.text == body]
        assert matching, body
        assert all(p.font.name == "Liberation Mono" and p.font.size.pt == 22 for shape in matching for p in shape.text_frame.paragraphs)
    assert all(record["source_refs"] for record in records)
    for record in records:
        assert record["source_files"] == {"teacher": str(session.teacher_source), "student_fallback": str(session.student_source)}
        assert all(ref["source"] in record["source_files"] for ref in record["source_refs"])
        assert all(need["status"] == "resolved" for need in record["representation_needs"])
    assert sum(record["timing"]["seconds"] for record in records) == 2700


@pytest.mark.parametrize("number", [str(number) for number in range(206, 216)])
def test_semantic_shapes_stay_in_bounds_with_large_fonts_and_no_shrinking(generator, tmp_path, number):
    from pptx.enum.text import MSO_AUTO_SIZE

    session = generator.parse_session(*source_pair(number))
    target = tmp_path / f"S{number}.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    for index, slide in enumerate(deck.slides):
        for shape in slide.shapes:
            assert 0 <= shape.left and shape.left + shape.width <= deck.slide_width
            assert 0 <= shape.top and shape.top + shape.height <= deck.slide_height, (number, index, shape.name)
            if shape.has_text_frame and shape.text:
                assert shape.text_frame.auto_size == MSO_AUTO_SIZE.NONE
                for p in shape.text_frame.paragraphs:
                    assert p.font.size.pt >= (16 if shape.name == "Sesión y procedencia" else 20)
        content = [shape for shape in slide.shapes if shape.has_text_frame and shape.text and shape.name not in {
            "Código de sesión", "Título", "Fase HEXA", "Modalidad", "Contexto",
            "Idea principal", "Sesión y procedencia",
        }]
        assert all(shape.top + shape.height <= Inches(6.84) for shape in content), [
            (number, index + 1, shape.name, (shape.top + shape.height) / Inches(1))
            for shape in content if shape.top + shape.height > Inches(6.84)
        ]
        titles = [shape for shape in slide.shapes if shape.name == "Idea principal"]
        if titles and content:
            assert titles[0].top + titles[0].height <= min(shape.top for shape in content), (
                number, index + 1, titles[0].text,
            )
        contents = [shape for shape in slide.shapes if shape.has_text_frame and shape.name not in
                    {"Código de sesión", "Título", "Fase HEXA", "Modalidad", "Contexto", "Idea principal", "Sesión y procedencia"}]
        for a in contents:
            for b in contents:
                if a is b or a.name == "Destino de clasificación" or b.name == "Destino de clasificación":
                    continue
                overlap = a.left < b.left + b.width and b.left < a.left + a.width and a.top < b.top + b.height and b.top < a.top + a.height
                assert not overlap, (number, index, a.name, b.name)


def test_relations_and_pilot_curriculum_survive_in_editable_projection(generator, tmp_path):
    expected = {
        "206": ["Alcance", "Correcto", "Eficiente", "Mantenible", "Scrum"],
        "212": ["Scanner", "nextLine", "Integer.parseInt", "Double.parseDouble", "(int) valor", "3.999", "no redondea", "NumberFormatException", "Boolean.parseBoolean"],
        "213": ["Limpieza, nombres claros y simplicidad", "nombres significativos", "comentarios útiles", "simplicidad", "incremento H1"],
        "214": ["README", "Entrada", "esperad", "obtenid", "Demuestra", "permiso", "externa", "Moodle"],
        "215": ["Defensa", "recuperación", "parejas", "Review", "Retrospectiva", "Moodle"],
    }
    special = set()
    for number, phrases in expected.items():
        session = generator.parse_session(*source_pair(number))
        target = tmp_path / f"S{number}.pptx"
        generator.build_presentation(session, target)
        deck = Presentation(target)
        text = "\n".join(visible_text(slide) for slide in deck.slides).lower()
        for phrase in phrases:
            assert phrase.lower() in text, (number, phrase)
        special.update(metadata(slide)["layout"] for slide in deck.slides)
        if number in {"212", "213"}:
            last = visible_text(deck.slides[-1])
            assert generator.plan_slides(session)[-1].visible_content[-1].lstrip("> ").replace("`", "") in last
    assert {"prediction", "recognition", "parallel", "check", "closure"} <= special
    assert {"contrast", "comparison"} & special


def test_s213_does_not_promote_decision_protocols_to_projected_core(generator, tmp_path):
    session = generator.parse_session(*source_pair("213"))
    target = tmp_path / "S213.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    visible = "\n".join(visible_text(slide) for slide in deck.slides).lower()
    assert all(term not in visible for term in ("decisiones anidadas", "operador ternario", "condiciones encadenadas", "menús"))
    assert not any(record == "core" for slide in deck.slides for unit, record in metadata(slide)["roles"].items()
                   if any(term in visible_text(slide).lower() for term in ("if/else", "ternario", "anidad")))


def test_s213_simple_expressions_are_recognition_not_core(generator, tmp_path):
    session = generator.parse_session(*source_pair("213"))
    target = tmp_path / "S213.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    recognition = [slide for slide in deck.slides if metadata(slide)["recognition"] and
                   "boolean tieneNombre" in visible_text(slide)]
    assert recognition
    assert all(set(metadata(slide)["roles"].values()) == {"recognition"} for slide in recognition)


def test_parallel_long_delivery_label_keeps_three_lines_at_24pt(generator, tmp_path):
    session = generator.parse_session(*source_pair("215"))
    target = tmp_path / "S215.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    slide = next(slide for slide in deck.slides if metadata(slide)["layout"] == "parallel")
    box = next(shape for shape in slide.shapes if shape.has_text_frame and "Entrega oficial Moodle" in shape.text)
    assert box.text_frame.paragraphs[0].font.size.pt == 24
    required = 3 * 24 * 1.05 + (box.text_frame.margin_top + box.text_frame.margin_bottom) / 12700
    assert box.height / 12700 >= required


def test_counterexample_has_textual_status_not_only_color(generator, tmp_path):
    session = generator.parse_session(*source_pair("209"))
    target = tmp_path / "S209.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    contrasts = [slide for slide in deck.slides if metadata(slide)["layout"] in {"contrast", "comparison"}]
    assert contrasts
    assert all("Versión problemática" in visible_text(slide) and "Alternativa" in visible_text(slide) for slide in contrasts)


def test_an_oversized_program_is_refused_whole_instead_of_shrunk_or_cut(generator):
    import render_presentaciones_semanticas as renderer

    program = "```java\npublic class Main {\n    public static void main(String[] argumentos) {\n" + "\n".join(
        f"        int dato{index} = {index};" for index in range(60)) + "\n    }\n}\n```"
    unit = generator.PedagogicalUnit("oversize", [generator.SourceRef("teacher", "Programa completo", 0)], "example", "core", [program])
    spec = generator.SlideSpec("focus", "Programa completo", visible_content=[program])
    spec.pedagogical_units = [unit]
    spec.unit_dispositions = {unit.unit_id: generator.UnitDisposition("visible", "Programa íntegro", ("visible",))}
    session = generator.parse_session(*source_pair("212"))
    with pytest.raises(renderer.RepresentationError, match="oversize.*no se recorta"):
        renderer.render_semantic_presentation(Presentation(), session, [spec])


@pytest.mark.parametrize("number", ["203", "225", "267", "284", "306"])
def test_nonadopted_renderer_keeps_the_legacy_pptx(generator, tmp_path, number):
    from zipfile import ZipFile

    session = generator.parse_session(*source_pair(number))
    reference = Presentation()
    reference.slide_width, reference.slide_height = generator.SLIDE_W, generator.SLIDE_H
    for spec in generator.plan_slides_legacy(session):
        generator.RENDERERS[spec.kind](reference, session, spec)
    expected = tmp_path / "legacy.pptx"
    actual = tmp_path / "dispatch.pptx"
    reference.save(expected)
    generator.build_presentation(session, actual)
    with ZipFile(expected) as a, ZipFile(actual) as b:
        assert {name: a.read(name) for name in a.namelist()} == {name: b.read(name) for name in b.namelist()}


def test_recovery_projects_student_actions_not_internal_teacher_requests(generator, tmp_path):
    session = generator.parse_session(*source_pair("215"))
    target = tmp_path / "S215.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    recovery = [slide for slide in deck.slides if "Recuperación ante una defensa insuficiente" in visible_text(slide)]
    assert recovery
    projected = "\n".join(visible_text(slide) for slide in recovery)
    assert "pide señalar" not in projected and "solicita una predicción" not in projected
    assert "Señala la línea" in projected and "Predice con los valores indicados" in projected and "Vuelve a explicar" in projected
    assert all("INDIVIDUAL" in visible_text(slide) for slide in recovery)
    assert all(metadata(slide)["layout"] == "check" for slide in recovery)
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in recovery)
    assert "pide señalar la línea" in notes and "registra para el seguimiento docente" in notes


def test_s212_pptx_orders_individual_before_pairs_and_keeps_the_correct_clock_source(generator, tmp_path):
    session = generator.parse_session(*source_pair("212"))
    target = tmp_path / "S212.pptx"
    generator.build_presentation(session, target)
    deck = Presentation(target)
    individual = [i for i, slide in enumerate(deck.slides) if "Leer, convertir, predecir y probar" in visible_text(slide)]
    pairs = [i for i, slide in enumerate(deck.slides) if "Comparar, probar y explicar" in visible_text(slide)]
    assert max(individual) < min(pairs)
    assert all(metadata(deck.slides[i])["timeline_refs"] == [2] for i in individual)
    assert all(metadata(deck.slides[i])["timeline_refs"] == [3] for i in pairs)


@pytest.mark.parametrize("number", ["206", "212", "213", "214", "215"])
def test_content_panels_contain_explicit_line_boxes(generator, tmp_path, number):
    """A font multiplier is not a portable, measurable line height in Office."""
    import textwrap

    target = tmp_path / f"S{number}.pptx"
    generator.build_presentation(generator.parse_session(*source_pair(number)), target)
    for index, slide in enumerate(Presentation(target).slides, 1):
        for box in slide.shapes:
            if not box.has_text_frame or not box.text or box.name in {
                "Código de sesión", "Título", "Fase HEXA", "Modalidad", "Contexto",
                "Idea principal", "Sesión y procedencia", "Destino de clasificación", "Propuesta editable",
            }:
                continue
            tf = box.text_frame
            required = (tf.margin_top + tf.margin_bottom) / 12700
            for paragraph in tf.paragraphs:
                spacing = paragraph.line_spacing
                assert not isinstance(spacing, float), (number, index, box.name, "interlineado relativo")
                columns = max(1, int(((box.width - tf.margin_left - tf.margin_right) / 12700) /
                                     (paragraph.font.size.pt * 0.6)))
                lines = max(1, len(textwrap.wrap(paragraph.text, columns,
                                               replace_whitespace=False, drop_whitespace=False)))
                required += lines * spacing.pt
            assert box.height / 12700 >= required - 0.1, (number, index, box.name, required)


@pytest.mark.parametrize("number,slide_number", [("214", 16), ("215", 10)])
def test_cards_have_lateral_padding_and_a_clear_footer_zone(generator, tmp_path, number, slide_number):
    from pptx.util import Inches

    target = tmp_path / f"S{number}.pptx"
    generator.build_presentation(generator.parse_session(*source_pair(number)), target)
    marker = "Un enlace profundo lleva" if number == "214" else "solicitar una predicción antes de ejecutar"
    slide = next(s for s in Presentation(target).slides if marker in visible_text(s))
    contents = [s for s in slide.shapes if s.name in {"Contenido", "Código editable"}
                or s.name.startswith(("Acción o criterio ", "Soporte fuente"))]
    assert contents
    for box in contents:
        assert box.text_frame.margin_left >= Inches(0.12)
        assert box.text_frame.margin_right >= Inches(0.12)
        assert box.top + box.height <= Inches(6.85), (number, slide_number, box.name)


def test_quality_criteria_are_comparable_without_resolving_classification(generator, tmp_path):
    target = tmp_path / "S206.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("206")), target)
    deck = Presentation(target)
    comparison = [s for s in deck.slides if all(word in visible_text(s) for word in
                  ("Correcto", "Eficiente", "Mantenible"))]
    assert comparison
    assert "resolver el reto sin añadir complejidad innecesaria" in visible_text(comparison[0])
    assert "Tres criterios para juzgar el alcance" in visible_text(comparison[0])
    assert "vuestro H1" in visible_text(comparison[0])
    assert "Estas tres ideas sirven para juzgar el alcance" in comparison[0].notes_slide.notes_text_frame.text


def test_scope_clarification_remains_visible_and_traceable(generator, tmp_path):
    import hashlib

    session = generator.parse_session(*source_pair("206"))
    target = tmp_path / "S206.pptx"
    generator.build_presentation(session, target)
    slide = next(s for s in Presentation(target).slides if
                 "En el producto H1, el núcleo obligatorio debe aparecer integrado en Main.java." in visible_text(s))
    label = next(text for spec in generator.plan_slides(session) for text in spec.visible_content
                 if text.startswith("H1 recorre fundamentos suficientes"))
    assert any(source["sha256"] == hashlib.sha256(label.encode()).hexdigest()
               for source in metadata(slide)["notes_sources"])


def test_source_reserved_prediction_has_no_revelation_in_its_first_view(generator, tmp_path):
    target = tmp_path / "S212.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("212")), target)
    slides = list(Presentation(target).slides)
    prediction = next(i for i, s in enumerate(slides) if "¿Qué ocurrirá si Integer.parseInt" in visible_text(s))
    first = visible_text(slides[prediction])
    assert "Error de ejecución:" not in first
    assert "puede compilar y después fallar durante la ejecución" not in first
    assert metadata(slides[prediction])["relations"]["reveal_phase"] == "prediction"
    assert any("puede compilar y después fallar durante la ejecución" in visible_text(s)
               for s in slides[prediction + 1:])


def test_casting_prediction_precedes_its_explanation(generator, tmp_path):
    target = tmp_path / "S212.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("212")), target)
    slides = list(Presentation(target).slides)
    question = next(i for i, slide in enumerate(slides) if "¿El resultado será 3 o 4?" in visible_text(slide))
    explanation = next(i for i, slide in enumerate(slides) if "este casting no redondea" in visible_text(slide))
    assert question < explanation
    assert "no redondea" not in visible_text(slides[question])
    assert "double valor = 3.999;" in visible_text(slides[question])
    assert metadata(slides[question])["relations"]["reveal_phase"] == "prediction"
    assert metadata(slides[explanation])["relations"]["reveal_phase"] == "contrast"


def test_conceptual_arrows_are_not_rendered_as_java_code(generator, tmp_path):
    targets = {
        "212": ("String → parseo → número", "7 → 7.0", "precioEntero → 12", "notaEntera → 7"),
        "214": ("entrada → resultado esperado → resultado obtenido → qué demuestra",),
    }
    for number, expected in targets.items():
        target = tmp_path / f"S{number}.pptx"
        generator.build_presentation(generator.parse_session(*source_pair(number)), target)
        deck = Presentation(target)
        for relation in expected:
            shape = next(shape for slide in deck.slides for shape in slide.shapes
                         if shape.has_text_frame and relation in shape.text)
            if number != "214":
                assert shape.name != "Código editable"
            assert "->" not in shape.text

        assert not any(
            "->" in shape.text
            for slide in deck.slides
            for shape in slide.shapes
            if shape.has_text_frame and shape.name != "Código editable"
        )


def test_real_java_arrows_remain_code_during_conceptual_normalization():
    import enriquecer_soportes_semanticos as supports
    import render_presentaciones_semanticas as renderer

    conceptual = renderer.RenderAtom("dato -> resultado", (), "code", "text")
    markdown = renderer.RenderAtom("Caso A -> resultado.", (), "code", "markdown")
    java = renderer.RenderAtom("valores.stream().map(x -> x + 1);", (), "code", "java")
    normalized, changes = supports.normalize_presentation_atom(conceptual, set())
    normalized_markdown, markdown_changes = supports.normalize_presentation_atom(markdown, set())
    preserved, java_changes = supports.normalize_presentation_atom(java, set())
    assert normalized.kind != "code" and normalized.text == "dato → resultado"
    assert changes == {"dato -> resultado": "dato → resultado"}
    assert normalized_markdown.kind == "code" and normalized_markdown.language == "markdown"
    assert normalized_markdown.text == "Caso A → resultado."
    assert markdown_changes == {"Caso A -> resultado.": "Caso A → resultado."}
    assert preserved == java and java_changes == {}


def test_s211_renders_after_rejecting_stale_study_merge(generator, tmp_path):
    target = tmp_path / "S211.pptx"

    count = generator.build_presentation(generator.parse_session(*source_pair("211")), target)

    deck = Presentation(target)
    assert count == len(deck.slides)
    visible = "\n".join(visible_text(slide) for slide in deck.slides)
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    assert "int horasTotales = 3 + 2;" in visible
    assert "int dias = 5;\nint horasPorDia = 2;\nint horasTotales = dias * horasPorDia;" in visible
    assert "¿Qué expresión se evalúa, qué resultado produce, qué tipo tiene y dónde queda guardado?" in visible + notes
    assert "Pide predecir horasTotales antes de ejecutar" in visible + notes


def test_study_capacity_uses_post_normalization_atom_kind(generator):
    import componer_presentaciones_semanticas as composition
    import render_presentaciones_semanticas as renderer

    unit = generator.PedagogicalUnit(
        "unit-test", [generator.SourceRef("teacher", "Expresiones", 0)],
        "concept", "core", ["Expresiones"],
    )
    frame = renderer.RenderFrame(0, "Expresiones", "code", [unit], [
        renderer.RenderAtom(
            "3 + 2\nhoras * 60\nhoras >= 4\ntieneNombre && tieneObjetivo\n\"Hola, \" + nombreUsuario",
            (0,), "code", "java",
        ),
        renderer.RenderAtom(
            "3 + 2                       -> int\n"
            "5 / 2.0                     -> double\n"
            "horas >= 4                  -> boolean\n"
            "\"Hola, \" + nombreUsuario    -> String",
            (1,), "code", "text",
        ),
        renderer.RenderAtom(
            "Una expresión combina valores, variables u operadores y produce un resultado.",
            (2,),
        ),
        renderer.RenderAtom(
            "Antes de ejecutar, pregunta por el valor y el tipo esperado.",
            (3,),
        ),
    ])

    composed = composition.compose(renderer, None, [frame])

    assert len(composed) == 2
    assert {index for part in composed for atom in part.atoms for index in atom.source_indices} == {0, 1, 2, 3}
    assert all(composition.frame_start(renderer, part) + sum(
        composition.physical_height(renderer, atom, part.recognition) + .1 for atom in part.atoms
    ) <= 6.84 for part in composed)


def test_capacity_partition_preserves_pair_prefix_prose_and_every_source_index(generator):
    import componer_presentaciones_semanticas as composition
    import render_presentaciones_semanticas as renderer

    unit = generator.PedagogicalUnit(
        "unit-test", [generator.SourceRef("teacher", "Operar, almacenar y utilizar", 8)],
        "concept", "core", ["contenido"],
    )
    pair = renderer.RenderAtom(
        "int horasTotales = 3 + 2;", (6, 9), "pair", "java",
        "int dias = 5;\nint horasPorDia = 2;\nint horasTotales = dias * horasPorDia;",
    )
    question = renderer.RenderAtom(
        "¿Qué expresión se evalúa, qué resultado produce, qué tipo tiene y dónde queda guardado?", (7,),
    )
    support = renderer.RenderAtom(
        "Pide predecir horasTotales antes de ejecutar y explicar qué representa cada operando.", (10,),
    )
    frames = [
        renderer.RenderFrame(2, "Operar, almacenar y utilizar", "contrast", [unit], [pair, question, support]),
        renderer.RenderFrame(2, "Operar, almacenar y utilizar", "code", [unit], [
            renderer.RenderAtom("int minutos = 4 * 60;\nSystem.out.println(minutos);", (11,), "code", "java"),
            renderer.RenderAtom("literales -> expresión -> resultado -> variable -> salida", (13,), "code", "text"),
            renderer.RenderAtom("4 * 60;", (15,), "code", "java"),
        ]),
        renderer.RenderFrame(2, "Operar, almacenar y utilizar", "concept", [unit], [
            renderer.RenderAtom("¿Dónde se guardaría el resultado y para qué se utilizaría?", (16,)),
        ]),
    ]

    composed = composition.compose(renderer, None, frames)

    atoms = [atom for frame in composed for atom in frame.atoms]
    notes = "\n".join(text for frame in composed for _, _, text in frame.notes)
    assert any(atom.kind == "pair" and atom.text == pair.text and atom.alternative == pair.alternative for atom in atoms)
    assert set(index for atom in atoms for index in atom.source_indices) == {6, 7, 9, 10, 11, 13, 15, 16}
    for expected in (question.text, support.text):
        assert expected in "\n".join(atom.text for atom in atoms) or expected in notes


def test_metadata_code_inventory_uses_both_pair_sides_and_preserves_repeated_java():
    from types import SimpleNamespace
    import render_presentaciones_semanticas as renderer

    repeated = "int horas = 4;"
    spec = SimpleNamespace(visible_content=[
        f"```java\n{repeated}\n```",
        "```text\nhoras -> dato\n```",
        f"```java\n{repeated}\n```",
        "```java\nSystem.out.println(horas);\n```",
    ])
    frame = renderer.RenderFrame(0, "Pares", "comparison", [], [
        renderer.RenderAtom(repeated, (0, 1), "pair", "java", "horas -> dato"),
        renderer.RenderAtom("horas -> dato", (1, 3), "pair", "text", "System.out.println(horas);"),
        renderer.RenderAtom(repeated, (2,), "code", "java"),
    ])

    assert renderer.metadata_code_blocks(spec, frame) == [
        repeated, repeated, "System.out.println(horas);",
    ]


def test_final_geometry_validation_fails_closed_on_footer_or_title_overlap(generator):
    import render_presentaciones_semanticas as renderer

    session = generator.parse_session(*source_pair("208"))
    frame = renderer.RenderFrame(0, "Título breve", "concept", [])
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = renderer.semantic_base(deck, session, frame)
    renderer.text_box(slide, "Contenido", "Fuera", 0.7, 6.7, 11.9, 0.5)

    with pytest.raises(renderer.RepresentationError, match="6.84"):
        renderer.validate_semantic_slide_geometry(slide)


def test_study_layout_keeps_fixed_code_and_reference_fonts(generator):
    from pptx.enum.text import MSO_AUTO_SIZE
    import componer_presentaciones_semanticas as composition
    import render_presentaciones_semanticas as renderer

    session = generator.parse_session(*source_pair("211"))
    unit = generator.PedagogicalUnit(
        "unit-test", [generator.SourceRef("teacher", "Expresiones", 0)],
        "concept", "core", ["Expresiones"],
    )
    frame = renderer.RenderFrame(0, "Expresiones", "study", [unit], [
        renderer.RenderAtom("int horas = 4;", (0,), "code", "java"),
        renderer.RenderAtom("El resultado conserva su tipo.", (1,)),
    ])
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    slide = composition.render_study(renderer, deck, session, frame)

    code = next(shape for shape in slide.shapes if shape.name == "Código editable")
    reference = next(shape for shape in slide.shapes if shape.name == "Referente o acción")
    assert {paragraph.font.size.pt for paragraph in code.text_frame.paragraphs} == {22}
    assert {paragraph.font.size.pt for paragraph in reference.text_frame.paragraphs} == {24}
    assert code.text_frame.auto_size == reference.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert code.top + code.height <= Inches(6.84)
    assert reference.top + reference.height <= Inches(6.84)


def test_s211_content_boxes_stay_above_footer_without_shrinking(generator, tmp_path):
    from pptx.enum.text import MSO_AUTO_SIZE

    target = tmp_path / "S211.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("211")), target)
    slides = list(Presentation(target).slides)
    content = [shape for slide in slides for shape in slide.shapes
               if shape.has_text_frame and shape.name in {
                   "Código editable", "Referente o acción", "Contenido",
               }]

    assert content
    assert all(shape.top + shape.height <= Inches(6.84) for shape in content)
    assert all(shape.text_frame.auto_size == MSO_AUTO_SIZE.NONE for shape in content)
    assert any(paragraph.font.size.pt == 28 for shape in content if shape.name == "Contenido"
               for paragraph in shape.text_frame.paragraphs)


def test_s213_cleaning_examples_remain_complete_java_at_22pt(generator, tmp_path):
    target = tmp_path / "S213.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("213")), target)
    deck = Presentation(target)
    expected = "final int MAX_HORAS = 8;\nint horasEstudio = 5;\nSystem.out.println(horasEstudio + \"/\" + MAX_HORAS);"
    cases = next(shape for slide in deck.slides for shape in slide.shapes
                 if shape.has_text_frame and expected == shape.text)
    assert cases.name in {"Código editable", "Contraste 2"}
    assert all(paragraph.font.size.pt == 22 for paragraph in cases.text_frame.paragraphs)


def test_s213_renaming_example_preserves_the_observable_output(generator, tmp_path):
    target = tmp_path / "S213.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("213")), target)
    deck = Presentation(target)
    text = "\n".join(visible_text(slide) for slide in deck.slides)
    assert "final int MAX = 8;" in text and "final int MAX_HORAS = 8;" in text
    assert 'System.out.println(h + "/" + MAX);' in text
    assert 'System.out.println(horasEstudio + "/" + MAX_HORAS);' in text
    assert "Antes de editar, registra la salida" in text


def test_initial_readme_model_marks_tests_as_incomplete(generator, tmp_path):
    target = tmp_path / "S214.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("214")), target)
    slides = list(Presentation(target).slides)
    model = next(i for i, slide in enumerate(slides) if "## Pruebas" in visible_text(slide))
    text = visible_text(slides[model])
    assert "ESTRUCTURA INICIAL" in text
    assert all(word in text for word in ("entrada", "esperado", "obtenido", "significado"))
    assert any("Entrada usada:" in visible_text(slide) and "Resultado obtenido:" in visible_text(slide)
               for slide in slides[model + 1:])


def test_java_example_questions_distinguish_preparation_effect_and_output(generator, tmp_path):
    target = tmp_path / "S206.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("206")), target)
    slide = next(slide for slide in Presentation(target).slides if "NOMBRE_ASISTENTE" in visible_text(slide))
    text = visible_text(slide)
    assert "¿qué prepara cada línea, qué efecto tiene y cuál produce salida?" in text
    assert "¿qué comportamiento visible aporta cada línea?" not in text
    assert not any(line.strip().startswith("¿") and line.rstrip().endswith(";") for line in text.splitlines())


def test_scanner_contrast_is_labelled_before_its_conclusion(generator, tmp_path):
    target = tmp_path / "S212.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("212")), target)
    slide = next(s for s in Presentation(target).slides if
                 "new Scanner(System.in).nextLine()" in visible_text(s))
    text = visible_text(slide)
    assert "Reutiliza" in text and "Contraste" in text
    assert "oculta la reutilización" in text
    code = [shape for shape in slide.shapes if shape.name.startswith("Contraste ")]
    assert len(code) == 2
    assert all(shape.width >= Inches(11.5) for shape in code)


def test_scanner_explanation_cards_use_a_grid_wide_enough_for_every_line(generator, tmp_path):
    target = tmp_path / "S212.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("212")), target)
    slide = next(s for s in Presentation(target).slides if
                 "new Scanner(System.in) prepara la lectura desde la entrada estándar." in visible_text(s))
    cards = [shape for shape in slide.shapes if shape.name.startswith("Acción o criterio ")]
    assert cards
    assert len({shape.left for shape in cards}) == 2


def test_s213_cleaning_protocol_keeps_its_code_referent(generator, tmp_path):
    target = tmp_path / "S213.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("213")), target)
    text = "\n".join(visible_text(slide) for slide in Presentation(target).slides)
    assert "¿Qué versión permite explicar el dato sin buscar todas sus apariciones?" in text
    assert "int x = 4;" in text and "int horasEstudio = 4;" in text
    assert "No añadáis funcionalidades nuevas durante esta revisión." in text


def test_s211_java_pair_is_complete_when_capacity_requires_partition(generator, tmp_path):
    target = tmp_path / "S211.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("211")), target)
    slides = [s for s in Presentation(target).slides if
              metadata(s)["layout"] == "comparison" and "int horasTotales = 3 + 2;" in visible_text(s)]
    assert slides
    for slide in slides:
        code = [shape for shape in slide.shapes if shape.name.startswith("Contraste ")]
        assert len(code) == 2
        assert any("int dias = 5;" in shape.text and "int horasTotales = dias * horasPorDia;" in shape.text
                   for shape in code)
        assert all(shape.top + shape.height <= Inches(6.84) for shape in code)
        assert all(paragraph.font.size.pt == 22 for shape in code for paragraph in shape.text_frame.paragraphs)


def test_saved_notes_have_no_empty_labels_or_algorithm_explanations(generator, tmp_path):
    for number in ("206", "212", "213", "214", "215"):
        target = tmp_path / f"S{number}.pptx"
        generator.build_presentation(generator.parse_session(*source_pair(number)), target)
        for slide in Presentation(target).slides:
            notes = slide.notes_slide.notes_text_frame.text
            assert "desempate por orden" not in notes and "sin ancla" not in notes.lower()
            assert not re.search(r"(?m)^(?:Pregunta(?: principal| complementaria)?|Cierra en voz alta|Cuando proceda):\s*$", notes)


def test_deduplicated_notes_do_not_leave_orphan_labels(generator, tmp_path):
    expected = {
        "206": {
            16: ("Al terminar, cada persona y cada equipo deben poder responder:",
                 "Cierra con la transición:"),
        },
        "212": {
            8: ("La secuencia conceptual es:", "Aplícala a una lectura real:", "Antes de ejecutar, pregunta:"),
            9: ("La secuencia conceptual es:", "Aplícala a una lectura real:", "Antes de ejecutar, pregunta:"),
        },
        "214": {
            4: ("Antes de aceptar una formulación, pregunta:",),
            22: ("Modelo de conexión:", "Cuando proceda:"),
        },
    }
    for number, slides in expected.items():
        target = tmp_path / f"S{number}.pptx"
        generator.build_presentation(generator.parse_session(*source_pair(number)), target)
        deck = Presentation(target)
        for slide_number, labels in slides.items():
            notes = deck.slides[slide_number - 1].notes_slide.notes_text_frame.text
            assert all(label not in notes for label in labels), (number, slide_number, notes)
            if number == "214" and slide_number == 22:
                label = "Como reflexión breve, sin crear una tarea adicional:"
                all_notes = "\n".join(s.notes_slide.notes_text_frame.text for s in deck.slides)
                assert label in all_notes and "¿Qué evidencia de H1" in all_notes.split(label, 1)[1]


def test_closure_names_the_students_code_as_the_real_referent(generator, tmp_path):
    target = tmp_path / "S213.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("213")), target)
    slide = next(s for s in Presentation(target).slides if "¿Qué nombre, comentario o elemento simplificaste" in visible_text(s))
    assert "comportamiento sigue siendo el mismo" in visible_text(slide)


def test_brief_warning_stays_with_its_peer_protocol(generator, tmp_path):
    target = tmp_path / "S215.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("215")), target)
    slide = next(s for s in Presentation(target).slides if "No preparan respuestas idénticas" in visible_text(s))
    assert "solicitar una predicción antes de ejecutar" in visible_text(slide)
    assert "ADVERTENCIA" in visible_text(slide)


def test_mixed_check_preserves_a_secondary_recognition_clause(generator, tmp_path):
    target = tmp_path / "S212.pptx"
    generator.build_presentation(generator.parse_session(*source_pair("212")), target)
    slide = next(s for s in Presentation(target).slides if "cuándo necesitas parsear" in visible_text(s))
    core = [s for s in slide.shapes if s.has_text_frame and "cuándo necesitas parsear" in s.text]
    secondary = [run for shape in slide.shapes if shape.has_text_frame
                 for p in shape.text_frame.paragraphs for run in p.runs if "por qué Integer.parseInt" in run.text]
    assert core and secondary and "PARA RECONOCER" in visible_text(slide)
    assert all(run.font.size.pt == 24 for run in secondary)
    assert all(run.font.color.rgb == __import__("pptx").dml.color.RGBColor(255, 255, 255)
               for shape in core for paragraph in shape.text_frame.paragraphs for run in paragraph.runs)


@pytest.mark.parametrize("number", ["206", "212", "213", "214", "215"])
def test_presenter_exclusive_units_never_leak_to_visible_shapes(generator, tmp_path, number):
    session = generator.parse_session(*source_pair(number))
    plan = generator.plan_slides(session)
    target = tmp_path / f"S{number}.pptx"
    generator.build_presentation(session, target)
    text = normalized("\n".join(visible_text(s) for s in Presentation(target).slides))
    exclusive = [(u, spec.unit_dispositions[u.unit_id]) for spec in plan for u in spec.pedagogical_units
                 if set(spec.unit_dispositions[u.unit_id].channels) == {"presenter"}]
    assert exclusive
    for unit, disposition in exclusive:
        for item in unit.presenter_content + unit.visible_content:
            if len(normalized(item)) > 25:
                assert normalized(item) not in text, (unit.unit_id, unit.function, disposition, item)
