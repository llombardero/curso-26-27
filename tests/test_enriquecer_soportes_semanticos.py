"""Source-backed supports: exercise the real planner, atoms and geometry."""
import importlib
import pytest
from pptx import Presentation
from pptx.util import Inches
from test_generar_presentaciones_sesiones import load_module, source_pair


def prepared(number, enriched=True):
    g = load_module()
    r = importlib.import_module('render_presentaciones_semanticas')
    session = g.parse_session(*source_pair(number))
    frames = r.expand_semantic_frames(session, g.plan_slides(session))
    if enriched:
        try:
            helper = importlib.import_module('enriquecer_soportes_semanticos')
        except ModuleNotFoundError as error:
            if error.name != 'enriquecer_soportes_semanticos':
                raise
        else:
            frames = helper.enrich(r, session, frames)
    return r, session, frames


def text(frame):
    return '\n'.join(a.text + '\n' + a.alternative for a in frame.atoms)


def test_external_reader_and_no_oral_help_are_visible_with_steps():
    _, _, frames = prepared('214')
    related = [f for f in frames if 'Comprobación externa' in f.title]
    corpus = '\n'.join(text(frame) for frame in related)
    assert 'abrir el enlace' in corpus
    assert 'sin explicación oral previa' in corpus
    assert 'otra persona' in corpus


def test_reproducible_example_question_and_chain_share_view():
    _, _, frames = prepared('214')
    frame = next(f for f in frames if 'Entrada usada: Laura.' in text(f))
    assert '¿Qué demuestra exactamente esta prueba?' in text(frame)
    assert 'entrada → resultado esperado → resultado obtenido → qué demuestra' in text(frame)


def test_prepare_deliver_boundary_is_prominent_exact_source():
    _, _, frames = prepared('214')
    exact = 'Moodle: H1.9 prepara enlaces, permisos y documentación; H1.10 realiza la entrega oficial de H1.'
    assert any(a.text == exact and a.kind == 'boundary' for f in frames for a in f.atoms)


def test_scope_lists_have_explicit_source_markers_without_solving_proposals():
    _, _, before = prepared('206', False)
    _, _, after = prepared('206')
    scopes = [f for f in after if f.title == 'Alcance técnico de H1']
    assert any(any(a.kind == 'cards' for a in f.atoms) for f in scopes)
    assert any('En el producto H1, el núcleo obligatorio debe aparecer integrado en Main.java' in text(f).replace('`', '') for f in scopes)
    assert any('H1 recorre fundamentos suficientes' in note[2] for f in scopes for note in f.notes)
    assert any('Quedan fuera del alcance actual:' in text(f) for f in scopes)
    assert any('contenidos posteriores' in text(f) for f in scopes)
    assert [text(f) for f in before if f.kind == 'classification'] == [text(f) for f in after if f.kind == 'classification']


@pytest.mark.parametrize('number', ['206', '214', '215'])
def test_all_supports_render_preserve_notes_units_and_valid_plan_indices(number, tmp_path):
    import json
    from tiempos_presentaciones_semanticas import allocate_slide_timings
    r, session, before = prepared(number, False)
    helper = importlib.import_module('enriquecer_soportes_semanticos')
    after = helper.enrich(r, session, before)
    assert helper.enrich(r, session, after) == after
    assert {u.unit_id for f in before for u in f.units} == {u.unit_id for f in after for u in f.units}
    assert {note for f in before for note in f.notes} <= {note for f in after for note in f.notes}
    g = load_module()
    plan = g.plan_slides(session)
    for frame in after:
        assert all(0 <= i < len(plan[frame.plan_index].visible_content) for a in frame.atoms for i in a.source_indices)
        json.dumps(frame.relations)
        if len(frame.origin_slides) > 1:
            for original in before:
                if set(original.origin_slides) <= set(frame.origin_slides):
                    assert original.plan_index == frame.plan_index
    corpus = r.prose_text('\n'.join(text(f) + '\n' + '\n'.join(n[2] for n in f.notes) for f in after))
    for frame in before:
        for atom in frame.atoms:
            for line in r.prose_text(atom.text).splitlines():
                assert line in corpus, ('Contenido perdido', line)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    timings = allocate_slide_timings(session, after)
    assert sum(t.seconds for t in timings) == 2700
    for frame, timing in zip(after, timings):
        layouts = {**r.SEMANTIC_LAYOUTS, **helper.layouts(r)}
        slide = layouts[frame.kind](deck, session, frame)
        visible = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
        slide.notes_slide.notes_text_frame.text = r.presenter_notes(frame, timing, visible)
        if frame.kind in helper.layouts(r):
            assert all(s.top + s.height <= Inches(6.84) for s in slide.shapes
                       if s.name.startswith(('Soporte ', 'Acción o criterio ')))
    target = tmp_path / ('S' + number + '.pptx')
    deck.save(target)
    assert len(Presentation(target).slides) == len(after)


def test_readme_four_identities_share_one_legible_support():
    r, session, frames = prepared('214')
    headings = {'Qué hace', 'Qué límites tiene', 'Cómo se ejecuta', 'Qué pruebas demuestran que funciona'}
    supports = [f for f in frames if headings & {u.source_refs[0].heading for u in f.units} and f.kind != 'title']
    assert 1 <= len(supports) <= 2, 'La zona segura puede partir el soporte, no aislar cada criterio'
    assert headings <= {u.source_refs[0].heading for frame in supports for u in frame.units}
    corpus = '\n'.join(text(frame) for frame in supports)
    assert 'Cada prueba debe indicar la entrada' in corpus
    assert 'ejecutar el punto de entrada' in corpus
    assert len({origin for frame in supports for origin in frame.origin_slides}) >= 5
    helper = importlib.import_module('enriquecer_soportes_semanticos')
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for frame in supports:
        slide = ({**r.SEMANTIC_LAYOUTS, **helper.layouts(r)})[frame.kind](deck, session, frame)
        assert all(s.top + s.height <= Inches(6.84) for s in slide.shapes
                   if s.name.startswith(('Soporte ', 'Acción o criterio ')))
        assert all(p.font.size.pt >= 24 for s in slide.shapes
                   if s.name.startswith(('Soporte ', 'Acción o criterio '))
                   for p in s.text_frame.paragraphs)
