from pathlib import Path
import hashlib
import re
import shutil
import subprocess
import sys
from zipfile import ZipFile

from openpyxl import load_workbook
from pptx import Presentation
import pytest

from test_generar_presentaciones_sesiones import load_module, source_pair

ROOT = Path(__file__).resolve().parents[1]
ALUMNADO = ROOT / "01-ALUMNADO"
PROFESORADO = ROOT / "02-PROFESORADO"


def test_hitos_sin_formularios_redundantes():
    hitos = sorted((ALUMNADO / "02-HITOS").iterdir())
    assert len(hitos) == 9
    for hito in hitos:
        assert not (hito / "plantillas").exists()
        names = {p.name for p in hito.iterdir() if p.is_file()}
        assert names == {"README.md", "ENTREGA-DIGITAL.md", next(n for n in names if "ficha-alumnado" in n)}


def test_sesiones_sin_bloque_administrativo_automatico():
    files = list((ALUMNADO / "03-SESIONES").rglob("S*-alumnado.md"))
    assert len(files) == 106
    for path in files:
        text = path.read_text(encoding="utf-8")
        assert "## Registro breve" not in text
        assert "## Evidencia mínima antes de salir" not in text


def test_modelo_transversal_minimo():
    diario = load_workbook(
        PROFESORADO / "05-ECOSISTEMA-DIGITAL/PLANTILLAS-MAESTRAS/01-Diario-individual-MiniJarvis.xlsx",
        read_only=True,
    )
    scrum = load_workbook(
        PROFESORADO / "05-ECOSISTEMA-DIGITAL/PLANTILLAS-MAESTRAS/02-Scrum-equipo-MiniJarvis.xlsx",
        read_only=True,
    )
    assert diario.sheetnames == ["DIARIO", "INSTRUCCIONES"]
    assert diario["DIARIO"].max_column == 7
    assert scrum.sheetnames == [
        "EQUIPO_Y_ACUERDOS", "BACKLOG_OPERATIVO", "DECISIONES_Y_BLOQUEOS",
        "REVIEW", "RETROSPECTIVA", "IA_EQUIPO", "INSTRUCCIONES",
    ]


def test_drive_operativo_no_es_espejo():
    files = [p for p in (ROOT / "04-DRIVE-5-EQUIPOS").rglob("*") if p.is_file()]
    assert len(files) == 19
    assert not any(p.suffix.lower() in {".html", ".java", ".pptx"} for p in files)
    assert len([p for p in files if p.name == "Scrum-equipo-MiniJarvis.xlsx"]) == 5


def test_moodle_distribuye_sin_ejemplos_privados():
    package = ROOT / "05-PAQUETE-MOODLE"
    assert len(list((package / "03-TAREAS-HITOS").glob("*.html"))) == 10
    assert not any("laura" in p.name.lower() for p in package.rglob("*"))
    h0 = (PROFESORADO / "05-ECOSISTEMA-DIGITAL/TAREAS-MOODLE/H0-tarea-moodle.md").read_text(encoding="utf-8")
    assert "no exige GitHub, tag, Site personal, Site de equipo" in h0


def test_html_y_presentaciones_sin_derivados_huerfanos():
    markdown = {p.relative_to(ALUMNADO).with_suffix(".html") for p in ALUMNADO.rglob("*.md")}
    html_root = ROOT / "01-ALUMNADO-HTML"
    html = {p.relative_to(html_root) for p in html_root.rglob("*.html")}
    assert html == markdown | {Path("LEEME-ALUMNADO.html")}
    pptx = sorted((PROFESORADO / "03-PRESENTACIONES/POR-SESION").rglob("S*-presentacion.pptx"))
    assert len(pptx) == 106
    student_by_code = {
        re.search(r"S\d{3}", source.name).group(0): source
        for source in (ALUMNADO / "03-SESIONES").rglob("S*-alumnado.md")
    }
    for path in pptx:
        presentation = Presentation(path)
        assert len(presentation.slides) >= 7
        code = re.search(r"S\d{3}", path.name).group(0)
        source = student_by_code[code].read_text(encoding="utf-8")
        mode = re.search(r"\*\*Modalidad:\*\* (.+?)\.", source).group(1)
        slide_text = "\n".join(
            shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text")
        )
        assert mode in slide_text
        assert "MINIJARVIS · PROGRAMACIÓN + ENTORNOS DE DESARROLLO" not in slide_text


def test_zips_validos_y_sin_temporales():
    archives = sorted(ROOT.glob("Minijarvis-*.zip"))
    assert len(archives) == 7
    for archive in archives:
        with ZipFile(archive) as zf:
            assert zf.testzip() is None
            assert not any(re.search(r"(^|/)(__pycache__|\.obsidian)(/|$)|\.pyc$", name) for name in zf.namelist())


@pytest.fixture(scope="module")
def session_parser():
    return load_module()


def assert_explicit_grouping(session):
    assert session.grouping.strip(), f"S{session.number}: modalidad ausente tras parseo"
    assert session.field_sources["grouping"] == "teacher", f"S{session.number}: modalidad no declarada en la guía"
    modes = {mode.upper() for mode in re.findall(r"\b(?:individual|parejas|equipo)\b", session.grouping, re.I)}
    assert modes, f"S{session.number}: agrupamiento no comprensible"
    if session.folder == "h1":
        # La guía canónica desarrolla las fases; las fichas conservan un resumen legacy.
        required = {mode for blocks in session.semantic_blocks.values() for block in blocks
                    for mode in re.findall(r"\b(?:INDIVIDUAL|PAREJAS|EQUIPO)\b", block.heading)}
        assert required <= modes, f"S{session.number}: agrupamiento perdido: {required - modes}"
    return modes


def test_modalidad_explicita_y_coherente_en_106_parejas(session_parser):
    allowed = {
        "Individual", "Equipo", "Parejas",
        "Individual → puesta en común en equipo",
        "Equipo → comprobación individual",
        "Equipo → comprobación cruzada",
        "Equipo → defensa individual",
        "Individual → contraste por parejas",
    }
    counts = {value: 0 for value in allowed}
    students = sorted((ALUMNADO / "03-SESIONES").rglob("S*-alumnado.md"))
    assert len(students) == 106
    for student in students:
        number = re.search(r"S(\d{3})", student.name).group(1)
        teacher = PROFESORADO / "02-SESIONES" / student.relative_to(ALUMNADO / "03-SESIONES")
        teacher = teacher.with_name(teacher.name.replace("-alumnado.md", "-docente.md"))
        student_text = student.read_text(encoding="utf-8")

        match = re.search(r"\*\*Modalidad:\*\* (.+?)\.", student_text)
        assert match, student
        mode = match.group(1)
        assert mode in allowed
        session = session_parser.parse_session(teacher, student)
        assert_explicit_grouping(session)
        if session.folder != "h1":
            assert session.grouping == mode, number
        counts[mode] += 1
        if "→" in mode and student.parent.name != "h1":
            assert "## Organización del trabajo" in student_text
    assert counts == {
        "Individual": 23,
        "Equipo": 49,
        "Parejas": 2,
        "Individual → puesta en común en equipo": 17,
        "Equipo → comprobación individual": 11,
        "Equipo → comprobación cruzada": 1,
        "Equipo → defensa individual": 1,
        "Individual → contraste por parejas": 2,
    }


@pytest.mark.parametrize("number, expected", [
    ("206", "INDIVIDUAL → EQUIPO"),
    ("210", "INDIVIDUAL → EQUIPO → comprobación INDIVIDUAL"),
    ("213", "INDIVIDUAL → PAREJAS → comprobación INDIVIDUAL"),
    ("215", "Defensa INDIVIDUAL; ensayo y revisión por PAREJAS; review, retrospectiva y entrega en EQUIPO"),
])
def test_modalidad_canonica_h1_conserva_sus_fases(session_parser, number, expected):
    session = session_parser.parse_session(*source_pair(number))
    assert_explicit_grouping(session)
    assert session.grouping == expected


@pytest.mark.parametrize("label", ["Modalidad", "Modalidad de trabajo", "Modalidad combinada", "Agrupamiento"])
def test_contrato_modalidad_admite_etiquetas_del_parser(session_parser, tmp_path, label):
    teacher, student = source_pair("206")
    text = teacher.read_text(encoding="utf-8").replace("| Modalidad de trabajo |", f"| {label} |")
    copy = tmp_path / "h1" / teacher.name
    copy.parent.mkdir()
    copy.write_text(text, encoding="utf-8")
    session = session_parser.parse_session(copy, student)
    assert assert_explicit_grouping(session) == {"INDIVIDUAL", "EQUIPO"}


@pytest.mark.parametrize("row", ["", "| Modalidad de trabajo | |", "| Modalidad de trabajo | ** ** |"])
def test_contrato_modalidad_rechaza_declaracion_ausente_o_vacia(session_parser, tmp_path, row):
    teacher, student = source_pair("206")
    text = teacher.read_text(encoding="utf-8").replace("| Modalidad de trabajo | **INDIVIDUAL → EQUIPO** |", row)
    copy = tmp_path / "h1" / teacher.name
    copy.parent.mkdir()
    copy.write_text(text, encoding="utf-8")
    session = session_parser.parse_session(copy, student)
    with pytest.raises(AssertionError, match="modalidad ausente|modalidad no declarada"):
        assert_explicit_grouping(session)


@pytest.mark.parametrize("grouping", ["", "INDIVIDUAL", "sin agrupamiento definido"])
def test_contrato_modalidad_detecta_perdida_en_el_modelo(session_parser, grouping):
    from dataclasses import replace

    session = session_parser.parse_session(*source_pair("213"))
    with pytest.raises(AssertionError, match="modalidad ausente|agrupamiento perdido|agrupamiento no comprensible"):
        assert_explicit_grouping(replace(session, grouping=grouping))


def test_modalidades_criticas_y_defensas():
    expected = {
        "S204": "Equipo → comprobación individual",
        "S215": "Equipo → defensa individual",
        "S240": "Equipo → comprobación individual",
        "S257": "Individual",
        "S274": "Equipo → comprobación individual",
        "S291": "Equipo → comprobación individual",
        "S296": "Equipo → comprobación individual",
        "S304": "Parejas",
        "S306": "Individual",
    }
    for path in (ALUMNADO / "03-SESIONES").rglob("S*-alumnado.md"):
        code = re.search(r"S\d{3}", path.name).group(0)
        if code in expected:
            assert f"**Modalidad:** {expected[code]}." in path.read_text(encoding="utf-8")


def test_laura_usa_fuentes_evolutivas_y_sin_docs_por_hito():
    laura = ROOT / "03-EJEMPLOS-LAURA-PRIVADOS"
    assert not list(laura.rglob("docs"))
    assert not list(laura.rglob("evidencias-digitales"))
    assert {p.name for p in (laura / "FUENTES-CURSO").iterdir()} == {
        "01-Diario-individual-MiniJarvis.xlsx",
        "02-Scrum-equipo-MiniJarvis.xlsx",
    }
    assert len(list((laura / "ENTREGAS-MOODLE").glob("*.md"))) == 9
    assert len(list((laura / "PORTFOLIOS-PERIODICOS").glob("*.md"))) == 6
    assert len(list(laura.rglob("*Site*"))) == 0
    for hito in [p for p in laura.iterdir() if p.is_dir() and p.name.startswith("h")]:
        assert (hito / "README.md").is_file()
        text = (hito / "README.md").read_text(encoding="utf-8")
        assert "## Defensa de Laura [INDIVIDUAL]" in text
        assert "Programación + Entornos" not in text


def test_entregas_laura_replican_el_contrato_moodle_minimo():
    examples = ROOT / "03-EJEMPLOS-LAURA-PRIVADOS/ENTREGAS-MOODLE"
    h0 = (examples / "H0-entrega.md").read_text(encoding="utf-8")
    assert "No se entrega GitHub, tag ni Site" in h0
    for number in range(1, 8):
        code = f"H{number}"
        text = (examples / f"{code}-entrega.md").read_text(encoding="utf-8")
        assert f"`h{number}-entrega`" in text
        assert "diario y Scrum actualizados" in text
        assert "Site personal" not in text
        assert "`PDF adjunto`" not in text
        assert "`XLSX adjunto`" not in text
    final = (examples / "HF-entrega.md").read_text(encoding="utf-8")
    assert "`hf-final`" in final and "Site personal final" in final and "Site de equipo final" in final


def test_generador_laura_reproduce_exactamente_la_salida_canonica(tmp_path):
    canonical = ROOT / "03-EJEMPLOS-LAURA-PRIVADOS"
    shutil.copy2(ROOT / "generar_evidencias_laura.py", tmp_path)
    shutil.copytree(canonical, tmp_path / canonical.name)

    subprocess.run(
        [sys.executable, "generar_evidencias_laura.py"],
        cwd=tmp_path,
        check=True,
        capture_output=True,
        text=True,
    )

    expected = {
        path.relative_to(canonical): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in canonical.rglob("*") if path.is_file()
    }
    generated = {
        path.relative_to(tmp_path / canonical.name): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in (tmp_path / canonical.name).rglob("*") if path.is_file()
    }
    assert len(generated) == 72
    assert generated == expected


def test_fuentes_sin_modelo_documental_obsoleto():
    roots = [
        ALUMNADO / "00-EMPIEZA-AQUI",
        ALUMNADO / "03-SESIONES",
        PROFESORADO / "01-GUIAS-POR-HITO",
        PROFESORADO / "02-SESIONES",
    ]
    forbidden = [
        "docs/portfolio-h", "plantillas/portfolio-h", "plantillas/defensa-h",
        "docs/registro-ia", "Plantilla de defensa H1",
        "documento/captura de ejecución",
    ]
    for base in roots:
        for path in base.rglob("*.md"):
            text = path.read_text(encoding="utf-8")
            for marker in forbidden:
                assert marker not in text, f"{marker}: {path}"


def test_guias_github_y_drive_declaran_fuentes_canonicas():
    start = ALUMNADO / "00-EMPIEZA-AQUI"
    github = (start / "15-introduccion-git-y-github-alumnado.md").read_text(encoding="utf-8")
    drive = (start / "14-guia-basica-drive-alumnado.md").read_text(encoding="utf-8")
    for phrase in ("código", "historial", "versiones", "README"):
        assert phrase in github
    assert "espacio operativo excepcional" in drive
    assert "02_EVIDENCIAS_NO_CODE" in drive
    assert "capturas de ejecución rutinarias" in drive
    assert "subcarpeta de hito solo cuando exista" in drive
